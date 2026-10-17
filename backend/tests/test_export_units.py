"""F07 — exportação PPTX/PDF: testes SEM navegador (sempre executados).

Cobrem pré-checagens (fonte aprovada, integridade, bloqueios), segurança (caminhos, nomes, argumentos do navegador,
requisições da página), fila/API (status, download, limites, expiração, auditoria), construtor PPTX com layout
sintético (texto, formas, custGeom, tabela/matriz, gráficos nativos com lacunas), fontes (WOFF→TTF→EOT), pacote
determinístico, validação de PPTX/PDF e regras de documento (matriz > 400 caracteres, rótulos "Crítica").
Dados 100% SINTÉTICOS; nenhuma Veracode.
"""
from __future__ import annotations

import hashlib
import io
import json
import struct
import time
import zipfile
from pathlib import Path

import pytest
from lxml import etree
from pptx import Presentation
from pptx.oxml.ns import qn
from sqlalchemy import text

from app import db
from app.export import browser, fonts as F, jobs as jobsmod, opc, pdf as pdfmod, pptx_builder, pptx_charts, render
from app.export import renderer, service as export_service, validate
from app.export.errors import ExportError
from app.export.service import ExportResult
from app.presentations import charts, export_plan
from conftest import error_code
from export_fixtures import SAST, full_values, ready_presentation
from test_presentations import OP, Env

# ---------------------------------------------------------------- utilidades

ORANGE = {'rgb': 'F26B1D', 'a': 1}
BLACK = {'rgb': '000000', 'a': 1}


@pytest.fixture
def env(api) -> Env:
    e = Env(api)
    e.builtin()
    return e


def run(**kw) -> dict:
    return {'t': 'Texto', 'font': 'Poppins', 'weight': 400, 'italic': False, 'size': 20, 'color': BLACK, 'spacing': 0,
            'underline': False, **kw}


def para(runs: list[dict], **kw) -> dict:
    return {'runs': runs, 'align': 'l', 'lh': 30, 'after': 0, 'bullet': None, 'marL': 0, 'indent': 0, **kw}


def wait_job(api, job: dict, timeout: float = 20) -> dict:
    end = time.monotonic() + timeout
    while job['status'] in ('queued', 'running') and time.monotonic() < end:
        time.sleep(0.05)
        job = api.get(f'/api/exports/{job["id"]}').json()
    return job


def tamper(db_url: str, *sql: str) -> None:
    engine = db.make_engine(db_url)
    try:
        with engine.begin() as conn:
            for s in sql:
                conn.execute(text(s))
    finally:
        engine.dispose()


# ---------------------------------------------------------------- pré-checagens (18, 19, 17)

def test_export_of_presentation_with_blockers_is_refused_with_diagnostics(env):
    c = env.client('Cliente Bloqueio F07')
    cid = env.approved(c['id'], '2026-09', {'apps.total': 10})
    p = env.present(cid, template_key='afrika-modelo', template_version=2)  # obrigatórios sem valor aprovado
    r = env.api.post(f'/api/presentations/{p["id"]}/exports', json={'format': 'pptx', 'revision': 1, 'operator': OP})
    assert r.status_code == 409 and error_code(r) == 'export_blocked'
    details = r.json()['error']['details']
    assert details and all(d['type'] == 'metric_missing' for d in details)
    assert all(d['field'].startswith('slides[') for d in details)  # diagnóstico por slide


def test_tampered_document_or_pinned_snapshot_blocks_export(env, db_url):
    c = env.client('Cliente Integridade F07')
    _, p, _ = ready_presentation(env, c['id'])
    tamper(db_url, 'DROP TRIGGER trg_metric_snapshots_update_draft',
           "UPDATE metric_snapshots SET value_json = '999' WHERE metric_id = 'apps.total'")
    r = env.api.post(f'/api/presentations/{p["id"]}/exports', json={'format': 'pdf', 'revision': 1, 'operator': OP})
    assert r.status_code == 409 and error_code(r) == 'export_blocked'
    assert any(d['type'] == 'integrity_failed' and 'apps.total' in d['message'] for d in r.json()['error']['details'])


def test_broken_source_reference_blocks_export(env, db_url):
    c = env.client('Cliente Referência F07')
    _, p, _ = ready_presentation(env, c['id'])
    tamper(db_url, 'DROP TRIGGER trg_presentation_revisions_no_update',
           "UPDATE presentation_revisions SET sources_json = replace(sources_json, 'APPROVED_ONLY_PINNED', 'DRAFT')")
    r = env.api.post(f'/api/presentations/{p["id"]}/exports', json={'format': 'pptx', 'revision': 1, 'operator': OP})
    assert r.status_code == 409 and error_code(r) == 'export_blocked'
    types = {d['type'] for d in r.json()['error']['details']}
    assert 'integrity_failed' in types  # fontes divergentes do hash registrado


def test_export_requires_explicit_persisted_revision_and_strict_body(env):
    c = env.client('Cliente Revisão F07')
    _, p, _ = ready_presentation(env, c['id'])
    base = f'/api/presentations/{p["id"]}/exports'
    assert env.api.post(base, json={'format': 'pptx', 'operator': OP}).status_code == 422  # sem revisão explícita
    r = env.api.post(base, json={'format': 'pptx', 'revision': 9, 'operator': OP})
    assert r.status_code == 404 and error_code(r) == 'presentation_revision_not_found'
    # o cliente não escolhe caminho/nome de arquivo: campos extras são recusados
    r = env.api.post(base, json={'format': 'pptx', 'revision': 1, 'operator': OP, 'path': '../../x.pptx'})
    assert r.status_code == 422
    assert env.api.post(base, json={'format': 'docx', 'revision': 1, 'operator': OP}).status_code == 422


def test_precheck_reads_the_requested_revision_not_a_draft(env):
    c = env.client('Cliente Rascunho F07')
    _, p, _ = ready_presentation(env, c['id'])
    doc = env.document(p['id'])
    doc['title'] = 'Título da revisão 2'
    env.save(p['id'], doc, 1)
    from app.export import service as svc
    from app.presentations import service as ps
    with env.api.app.state.session_factory() as s:
        pres = ps.get(s, __import__('uuid').UUID(p['id']))
        v1 = svc.precheck(s, pres, ps.get_revision(s, pres, 1))
        v2 = svc.precheck(s, pres, ps.get_revision(s, pres, 2))
    assert v1['presentation']['title'] != 'Título da revisão 2' and v1['revision']['number'] == 1
    assert v2['presentation']['title'] == 'Título da revisão 2' and v2['revision']['number'] == 2


# ---------------------------------------------------------------- fila, API, download, auditoria (21, 22, 25)

def fake_result(view: dict, fmt: str) -> ExportResult:
    data = (b'%PDF-1.4 sintetico' if fmt == 'pdf' else b'PK sintetico') + view['revision']['revision_hash'].encode()
    return ExportResult(data=data, filename=export_service.filename(view, fmt), media_type=export_service.FORMATS[fmt][0],
                        sha256=hashlib.sha256(data).hexdigest(), report={'slides': len(view['slides'])})


def test_job_lifecycle_download_audit_and_snapshots_untouched(env, db_url, tmp_path):
    c = env.client('Cliente Fila F07 / Ação & Cia')
    _, p, view = ready_presentation(env, c['id'])
    app = env.api.app
    app.state.exports.close()
    app.state.exports = jobsmod.ExportJobs(app.state.session_factory, base_dir=tmp_path / 'exp', runner=fake_result)
    engine = db.make_engine(db_url)
    with engine.connect() as conn:
        before = conn.execute(text('SELECT metric_id, row_hash FROM metric_snapshots ORDER BY id')).all()
    r = env.api.post(f'/api/presentations/{p["id"]}/exports', json={'format': 'pptx', 'revision': 1, 'operator': OP})
    assert r.status_code == 202
    job = wait_job(env.api, r.json())
    assert job['status'] == 'done' and job['download_url'] == f'/api/exports/{job["id"]}/file'
    assert 'path' not in job and 'owner_key' not in job  # caminho interno nunca exposto
    f = env.api.get(job['download_url'])
    assert f.status_code == 200 and f.content.startswith(b'PK')
    assert f.headers['x-content-type-options'] == 'nosniff' and f.headers['cache-control'] == 'no-store'
    disposition = f.headers['content-disposition']
    assert 'attachment' in disposition and '..' not in disposition and '/' not in job['filename']
    assert job['filename'].startswith('cliente-fila-f07-acao-cia_2026-09_') and job['filename'].endswith('_r1.pptx')
    listed = env.api.get(f'/api/presentations/{p["id"]}/exports').json()['items']
    assert [j['id'] for j in listed] == [job['id']]
    with engine.connect() as conn:
        after = conn.execute(text('SELECT metric_id, row_hash FROM metric_snapshots ORDER BY id')).all()
        audit = conn.execute(text("SELECT details_json, actor FROM audit_events WHERE event_type = "
                                  "'presentation.exported'")).all()
    engine.dispose()
    assert before == after  # exportar não altera snapshots
    assert len(audit) == 1 and audit[0][1] == OP
    assert json.loads(audit[0][0])['sha256'] == job['sha256'] and json.loads(audit[0][0])['format'] == 'pptx'
    assert env.call('GET', f'/api/presentations/{p["id"]}/integrity')['ok'] is True


def test_conversion_error_is_reported_and_no_file_is_served(env, tmp_path):
    c = env.client('Cliente Erro F07')
    _, p, _ = ready_presentation(env, c['id'])

    def broken(view, fmt):
        raise ExportError('export_engine_failed', 'O navegador de renderização encerrou ao iniciar', status=503)

    app = env.api.app
    app.state.exports.close()
    app.state.exports = jobsmod.ExportJobs(app.state.session_factory, base_dir=tmp_path / 'exp', runner=broken)
    job = wait_job(env.api, env.api.post(f'/api/presentations/{p["id"]}/exports',
                                         json={'format': 'pdf', 'revision': 1, 'operator': OP}).json())
    assert job['status'] == 'failed' and job['error']['code'] == 'export_engine_failed'
    assert job['download_url'] is None and job['sha256'] is None
    r = env.api.get(f'/api/exports/{job["id"]}/file')
    assert r.status_code == 409 and error_code(r) == 'export_not_ready'
    assert env.api.get('/api/exports/00000000-0000-4000-8000-000000000000').status_code == 404
    bad = env.api.get('/api/exports/..%2F..%2Fetc')  # id não-UUID nunca vira caminho (rota/validação recusam)
    assert bad.status_code in (404, 422) and bad.headers['content-type'].startswith('application/json')


def test_queue_limit_ttl_cleanup_and_shutdown(tmp_path):
    gate = __import__('threading').Event()

    def slow(view, fmt):
        gate.wait(5)
        return fake_result(view, fmt)

    view = {'presentation': {'id': 'p1', 'title': 'T'}, 'revision': {'number': 1, 'revision_hash': 'h'},
            'owner': {'key': 'client:1', 'name': 'Cliente'}, 'competence': {'competence': '2026-09'},
            'template': {'key': 't', 'version': 1}, 'theme': {'id': 'x', 'version': 1}, 'slides': []}
    q = jobsmod.ExportJobs(None, base_dir=tmp_path / 'q', runner=slow, max_queue=1, ttl=0.2)
    q._audit = lambda job: None
    first = q.submit(view, 'pptx', OP)
    time.sleep(0.2)  # o primeiro já está em execução; o segundo ocupa a fila
    q.submit(view, 'pdf', OP)
    with pytest.raises(ExportError) as exc:
        q.submit(view, 'pdf', OP)
    assert exc.value.code == 'export_queue_full' and exc.value.status == 429
    gate.set()
    end = time.monotonic() + 5
    while q.get(first.id).status != 'done' and time.monotonic() < end:
        time.sleep(0.05)
    path = q.get(first.id).path
    assert path.is_file() and path.is_relative_to(q.base)
    time.sleep(0.5)
    while any(j.status != 'done' for j in q.jobs.values()) and time.monotonic() < end + 5:
        time.sleep(0.05)
    time.sleep(0.3)
    with pytest.raises(ExportError):
        q.get(first.id)  # expirado => removido
    assert not path.exists()
    q.close()
    assert not q.base.exists()


def test_export_engine_endpoint_reports_without_launching(api, monkeypatch):
    monkeypatch.setenv(browser.BROWSER_ENV, 'relativo/chrome.exe')
    body = api.get('/api/export-engine').json()
    assert body['available'] is False and body['browser']['code'] == 'export_engine_unavailable'
    assert set(body['formats']) == {'pptx', 'pdf'}


# ---------------------------------------------------------------- segurança (25)

def test_filenames_are_generated_safely():
    v = {'owner': {'name': '../../Cliente "Ação"\\ Ltda.'}, 'competence': {'competence': '2026-09'},
         'presentation': {'title': 'Relatório: Setembro/2026 <script>'}, 'revision': {'number': 3}}
    name = export_service.filename(v, 'pptx')
    assert name == 'cliente-acao-ltda_2026-09_relatorio-setembro-2026-script_r3.pptx'
    assert all(ch.isalnum() or ch in '-_.' for ch in name)


def test_browser_launch_args_are_fixed_and_isolated(tmp_path, monkeypatch):
    exe = browser.BrowserExe(Path('/x/chrome'), 'system')
    args = browser._args(exe, tmp_path / 'perfil')
    assert args[0] == str(Path('/x/chrome')) and f'--user-data-dir={tmp_path / "perfil"}' in args
    assert '--remote-debugging-address=127.0.0.1' in args and '--host-resolver-rules=MAP * ~NOTFOUND' in args
    assert args[-1] == 'about:blank'
    monkeypatch.setenv(browser.BROWSER_ENV, str(tmp_path / 'inexistente.exe'))
    with pytest.raises(ExportError) as exc:
        browser.discover()
    assert exc.value.code == 'export_engine_unavailable'


class _FakeSession:
    def __init__(self):
        self.sent = []

    def send(self, method, params=None, **kw):
        self.sent.append((method, params))


def test_render_server_serves_only_dist_view_and_theme_assets(tmp_path):
    dist = tmp_path / 'dist'
    (dist / 'assets').mkdir(parents=True)
    (dist / 'index.html').write_text('<html></html>', encoding='utf-8')
    (dist / 'assets' / 'app.js').write_text('1', encoding='utf-8')
    (tmp_path / 'segredo.txt').write_text('não servir', encoding='utf-8')
    s = _FakeSession()
    pid = '11111111-1111-4111-8111-111111111111'
    srv = render._Server(s, renderer.Renderer(dist.resolve(), 'h', 'match'), pid, 2, b'{"view":1}', 'afrika',
                         {'logo-capa'})

    def ask(url, method='GET'):
        s.sent.clear()
        srv.on_request({'params': {'requestId': 'r', 'request': {'url': url, 'method': method}}})
        return s.sent[-1]

    assert ask(f'{render.ORIGIN}/index.html')[0] == 'Fetch.fulfillRequest'
    assert ask(f'{render.ORIGIN}/assets/app.js')[0] == 'Fetch.fulfillRequest'
    view = ask(f'{render.ORIGIN}/api/presentations/{pid}/view?revision=2')
    assert view[0] == 'Fetch.fulfillRequest'
    for url in (f'{render.ORIGIN}/../segredo.txt', f'{render.ORIGIN}/assets/%2e%2e/%2e%2e/segredo.txt',
                f'{render.ORIGIN}/api/presentations/{pid}/view?revision=1',  # outra revisão: bloqueada
                f'{render.ORIGIN}/api/presentations/22222222-2222-4222-8222-222222222222/view?revision=2',
                f'{render.ORIGIN}/api/clients', 'http://example.com/x.js', 'file:///etc/passwd',
                f'{render.ORIGIN}/api/brand-assets/outro-tema/logo-capa'):
        assert ask(url)[0] == 'Fetch.failRequest', url
    assert ask(f'{render.ORIGIN}/index.html', method='POST')[0] == 'Fetch.failRequest'
    assert srv.blocked and all('?' not in b for b in srv.blocked)  # registrado sem query string


def test_renderer_requires_fresh_build(tmp_path, monkeypatch):
    fe = tmp_path / 'frontend'
    (fe / 'src').mkdir(parents=True)
    (fe / 'src' / 'a.ts').write_text('export const a = 1;\n', encoding='utf-8')
    monkeypatch.setattr(renderer, 'REPO_ROOT', tmp_path)
    with pytest.raises(ExportError) as exc:
        renderer.locate({})
    assert exc.value.code == 'renderer_missing'
    (fe / 'dist').mkdir()
    (fe / 'dist' / 'index.html').write_text('x', encoding='utf-8')
    good, _ = renderer.source_hash(fe)
    (fe / 'dist' / renderer.STAMP_FILE).write_text(json.dumps({'source_hash': good}), encoding='utf-8')
    assert renderer.locate({}).stamp_check == 'match'
    (fe / 'src' / 'a.ts').write_text('export const a = 2;\n', encoding='utf-8')  # código mudou, build não
    with pytest.raises(ExportError) as exc:
        renderer.locate({})
    assert exc.value.code == 'renderer_stale'


# ---------------------------------------------------------------- overflow (13) e fontes ausentes (23)

def test_overflow_policy_blocks_real_cuts_and_spills_but_not_padding_intrusion():
    findings = [{'slide': '07-x', 'key': 'zone:detail', 'kind': 'clipped-y', 'amount': 8, 'cut': 0},
                {'slide': '14-matriz', 'key': 'matrix', 'kind': 'clipped-y', 'amount': 240, 'cut': 228},
                {'slide': '03-y', 'key': 'title', 'kind': 'spill-x', 'amount': 12, 'cut': 12},
                {'slide': '05-z', 'key': 'kpi', 'kind': 'outside-zone', 'amount': 3, 'cut': 3}]
    blocking, warnings = export_service.overflow_diagnostics(findings)
    assert [b['slide'] for b in blocking] == ['14-matriz', '03-y', '05-z']
    assert [w['element'] for w in warnings] == ['zone:detail']


def test_font_fallback_detection_uses_declared_web_fonts():
    lay = {'index': 2, 'slide_id': 's', 'ops': [
        {'op': 'text', 'key': 'text:title', 'xid': 1, 'paras': [para([run(font='Montserrat')])]},
        {'op': 'text', 'key': 'text:sys', 'xid': 2, 'paras': [para([run(font='Segoe UI')])]}]}
    cap = render.Capture([lay], [], {2: {'1': [{'family': 'Arial', 'custom': False, 'glyphs': 10}],
                                         '2': [{'family': 'Segoe UI', 'custom': False, 'glyphs': 4}]}},
                         webfonts=[], pdf=None, engine={}, declared_fonts=['Montserrat'])
    out = export_service.font_fallbacks(cap)
    assert [f['element'] for f in out] == ['text:title'] and out[0]['rendered_with'] == ['Arial']


# ---------------------------------------------------------------- construtor PPTX com layout sintético

def synthetic_view() -> dict:
    theme = {'id': 'afrika', 'version': 2, 'definition_hash': 'x' * 64, 'canvas': {'width': 1920, 'height': 1080},
             'colors': {'text.primary': '#2B0B43', 'surface': '#FFFFFF', 'band': '#F2EEF5', 'grid': '#F2EEF5',
                        'axis': '#DCC6EB', 'text.muted': '#888888', 'series.1': '#5D2482', 'series.2': '#2B0B43',
                        'series.3': '#D72369', 'series.4': '#FE6041', 'accent': '#5D2482', 'title': '#4A1D6E',
                        'severity.VERY_HIGH': '#C00000', 'severity.HIGH': '#F26B1D', 'severity.UNKNOWN': '#B5B5B5'},
             'fonts': {r: {'stack': [f, 'Arial'], 'export_family': f} for r, f in
                       (('heading', 'Montserrat'), ('body', 'Poppins'), ('numeric', 'Be Vietnam Pro'))},
             'weights': {'title': 900, 'value_label': 600, 'emphasis': 700},
             'sizes': {'axis': 16, 'value_label': 16, 'severity': 18.67},
             'charts': {'plot_padding': {'top': 28, 'right': 24, 'bottom': 48, 'left': 56}, 'plot_width': 900,
                        'bar_max_width': 48, 'bar_width_ratio': 0.6, 'y_ticks': 4, 'grid_width': 1, 'line_width': 3,
                        'point_radius': 6, 'point_stroke': 2, 'severity_bar_height': 28},
             'components': {'donut_hole': 0.6, 'donut_separator': 3}}
    pts = [{'x': m, 'y': y, 'state': st, 'label': lab} for m, y, st, lab in
           (('2026-07', None, 'GAP', 'lacuna'), ('2026-08', 139, 'VALUE', '139'), ('2026-09', 140, 'VALUE', '140'))]
    pts2 = [{'x': m, 'y': y, 'state': st, 'label': lab} for m, y, st, lab in
            (('2026-07', None, 'GAP', 'lacuna'), ('2026-08', 20, 'VALUE', '20'), ('2026-09', 1234, 'VALUE', '1.234'))]
    combo = {'kind': 'combo', 'available': True, 'categories': ['2026-07', '2026-08', '2026-09'],
             'series': [{'label': 'Aplicações', 'color': 'series.1', 'points': pts},
                        {'label': 'Findings', 'color': 'series.2', 'points': pts2}],
             'y_axis': {'integer': True, 'label': 'aplicações'}, 'number_format': {'decimals': 0, 'suffix': ''},
             'gaps': ['2026-07'], 'data_refs': [], 'show_values': True}
    sev = {'kind': 'severity', 'available': True, 'categories': ['Crítica', 'Alta', 'Desconhecida'],
           'series': [{'label': 'SAST', 'color': 'series.1', 'points': [
               {'x': 'Crítica', 'y': 2, 'state': 'VALUE', 'label': '2', 'color': 'severity.VERY_HIGH',
                'key': 'VERY_HIGH', 'source_label': 'Muito alta'},
               {'x': 'Alta', 'y': 0, 'state': 'VALUE', 'label': '0', 'color': 'severity.HIGH'},
               {'x': 'Desconhecida', 'y': 1, 'state': 'VALUE', 'label': '1', 'color': 'severity.UNKNOWN'}]}],
           'total': {'y': 3, 'label': '3'}, 'data_refs': [], 'gaps': []}
    slides = [{'index': 1, 'slide_id': 's1', 'type': 'dashboard', 'layout': 'painel-graficos',
               'fields': {'notes': 'Nota sintética do apresentador'},
               'data': {'kind': 'dashboard', 'zones': {'chart': {'data': combo}, 'donut_a': {'data': sev},
                                                       'sev': {'data': sev}}}}]
    return {'presentation': {'id': '11111111-1111-4111-8111-111111111111', 'title': 'Sintético'},
            'revision': {'number': 1, 'created_at': '2026-10-16T12:00:00+00:00', 'revision_hash': 'r' * 64,
                         'document_hash': 'd' * 64, 'sources_hash': 's' * 64},
            'owner': {'name': 'Cliente', 'kind': 'client', 'id': 'c1', 'key': 'client:c1'},
            'competence': {'competence': '2026-09'}, 'template': {'key': 'afrika-modelo', 'version': 2},
            'theme': theme, 'sources': {'primary': {'competence': '2026-09', 'revision': 1, 'approval_sequence': 1},
                                        'history': {}}, 'slides': slides}


def synthetic_layout() -> list[dict]:
    long_cell = ('Ação sintética ' * 40).strip()  # > 400 caracteres numa célula da matriz (runs já normalizados)
    cell = lambda t, x, w: {'cell': {'x': x, 'y': 300, 'w': w, 'h': 200}, 'box': {'x': x + 16, 'y': 312, 'w': w - 32, 'h': 60},  # noqa: E731
                            'paras': [para([run(t=t)], lh=30)] if t else [], 'wrap': True, 'vertical': False,
                            'lines': 2, 'lineWidths': [], 'maxWidth': w - 32, 'base': 24,
                            'fill': {'rgb': 'F2EEF5', 'a': 1} if t else None, 'pad': [12, 16, 12, 16],
                            'borders': [None, None, None, None], 'valign': 't', 'xid': 9}
    ops = [
        {'op': 'background', 'key': 'background', 'fill': {'rgb': 'FFFFFF', 'a': 1}, 'image': None},
        {'op': 'path', 'key': 'path:shape:painel', 'box': {'x': 0, 'y': 0, 'w': 997, 'h': 1080},
         'd': 'M0 0L792 0C927 153 997 328 997 520A400 400 0 0 1 752 1080L0 1080Z', 'vw': 997, 'vh': 1080,
         'fill': {'rgb': 'FFFFFF', 'a': 0.95}},
        {'op': 'rect', 'key': 'rect:pill', 'box': {'x': 10, 'y': 10, 'w': 100, 'h': 30},
         'fill': {'rgb': '2E9E5B', 'a': 0.18}, 'radius': [15, 15, 15, 15], 'border': None, 'texture': False},
        {'op': 'rect', 'key': 'rect:card', 'box': {'x': 200, 'y': 10, 'w': 300, 'h': 200}, 'fill': None,
         'radius': [32, 32, 0, 0], 'border': {'w': 4, 'color': {'rgb': 'F2EEF5', 'a': 1}, 'style': 'dashed'},
         'texture': False},
        {'op': 'rect', 'key': 'rect:unknown', 'box': {'x': 600, 'y': 10, 'w': 8, 'h': 60}, 'fill': {'rgb': 'B5B5B5', 'a': 1},
         'radius': [0, 0, 0, 0], 'border': None, 'texture': True},
        {'op': 'line', 'key': 'line:sep', 'x1': 0, 'y1': 250, 'x2': 500, 'y2': 250, 'w': 2,
         'color': {'rgb': 'F2EEF5', 'a': 1}, 'style': 'solid'},
        {'op': 'arrow', 'key': 'arrow:cmp', 'box': {'x': 700, 'y': 20, 'w': 28, 'h': 14}, 'color': BLACK, 'w': 2.3},
        {'op': 'text', 'key': 'text:title', 'xid': 1, 'box': {'x': 700, 'y': 60, 'w': 400, 'h': 73},
         'paras': [para([run(t='Título', font='Montserrat', weight=900, size=66.67)], lh=73.34, align='ctr')],
         'wrap': False, 'vertical': False, 'lines': 1, 'lineWidths': [400], 'maxWidth': 1000, 'base': 60,
         'prov': 'apps.total|2026-09|VALUE|MANUAL'},
        {'op': 'text', 'key': 'list:zone-list', 'xid': 2, 'box': {'x': 135, 'y': 500, 'w': 756, 'h': 100},
         'paras': [para([run(t='Item um')], marL=30, indent=-20,
                        bullet={'char': '•', 'color': {'rgb': '5D2482', 'a': 1}, 'size': 18, 'font': 'Poppins',
                                'weight': 400, 'scale': 1.3}),
                   para([run(t='Item dois com\nquebra')])],
         'wrap': True, 'vertical': False, 'lines': 3, 'lineWidths': [], 'maxWidth': 726, 'base': 22},
        {'op': 'text', 'key': 'text:section', 'xid': 3, 'box': {'x': 200, 'y': 300, 'w': 128, 'h': 600},
         'paras': [para([run(t='Seção vertical', size=128)], lh=128)], 'wrap': False, 'vertical': True, 'lines': 1,
         'lineWidths': [], 'maxWidth': 1080, 'base': 100},
        {'op': 'table', 'key': 'table:matrix', 'box': {'x': 136, 'y': 300, 'w': 1732, 'h': 420}, 'kind': 'matrix',
         'cols': [326, 452, 452], 'gap': 16, 'gapColor': {'rgb': 'FFFFFF', 'a': 1}, 'ir': None,
         'rows': [{'y': 300, 'h': 56, 'header': True, 'cells': [cell('', 136, 326), cell('Situação', 478, 452),
                                                               cell('Ação', 946, 452)]},
                  {'y': 372, 'h': 200, 'header': False, 'cells': [cell('Tema 1', 136, 326), cell(long_cell, 478, 452),
                                                                 cell('Curta', 946, 452)]}]},
        {'op': 'chart', 'key': 'chart:plot', 'box': {'x': 947, 'y': 766, 'w': 918, 'h': 264},
         'plot': {'x': 1003, 'y': 794, 'w': 838, 'h': 188}, 'ir': 'data.zones.chart.data', 'variant': 'plot', 'rows': 0,
         'vb': [918, 264], 'categories': ['2026-07', '', '2026-09'], 'zone': 'zone:chart',
         'overlays': [{'op': 'rect', 'key': 'rect:gap', 'box': {'x': 1020, 'y': 900, 'w': 48, 'h': 80}, 'fill': None,
                       'radius': [0, 0, 0, 0], 'border': {'w': 1, 'color': {'rgb': '888888', 'a': 1}, 'style': 'dashed'},
                       'texture': False}]},
        {'op': 'chart', 'key': 'chart:donut', 'box': {'x': 931, 'y': 491, 'w': 220, 'h': 220}, 'plot': None,
         'ir': 'data.zones.donut_a.data', 'variant': 'donut', 'rows': 0, 'vb': [220, 220], 'categories': [],
         'zone': 'zone:donut_a', 'overlays': []},
        {'op': 'chart', 'key': 'chart:sev-rows', 'box': {'x': 100, 'y': 700, 'w': 700, 'h': 200},
         'plot': {'x': 336, 'y': 700, 'w': 350, 'h': 126}, 'ir': 'data.zones.sev.data', 'variant': 'severity',
         'rows': 3, 'vb': [0, 0], 'categories': [], 'zone': None, 'overlays': []},
    ]
    return [{'version': 1, 'index': 1, 'slide_id': 's1', 'canvas': {'width': 1920, 'height': 1080}, 'ops': ops,
             'text_nodes': 3}]


def fake_book(root: Path) -> F.FontBook:
    """Catálogo de fontes com os NOMES de arquivo do build (o construtor só mapeia família/peso -> face)."""
    (root / 'assets').mkdir(parents=True, exist_ok=True)
    for fam, w in (('montserrat', 400), ('montserrat', 700), ('montserrat', 900), ('poppins', 400), ('poppins', 600),
                   ('poppins', 700), ('be-vietnam-pro', 400), ('be-vietnam-pro', 600), ('be-vietnam-pro', 700)):
        (root / 'assets' / f'{fam}-latin-{w}-normal-x.woff').write_bytes(b'x')
    return F.FontBook(root)


@pytest.fixture(scope='module')
def built(tmp_path_factory):
    book = fake_book(tmp_path_factory.mktemp('fontbook'))
    b = pptx_builder.build(pptx_builder.BuildInput(view=synthetic_view(), layouts=synthetic_layout(), fontbook=book))
    opc.add_custom_properties(b.prs, {'afrika.template': 'afrika-modelo@2', 'afrika.severity_labels': 'executive'})
    buf = io.BytesIO()
    b.prs.save(buf)
    data = opc.normalize(buf.getvalue(), pptx_builder._ts(synthetic_view()))
    return b, data


def shape(prs, name):
    return next(sh for sh in prs.slides[0].shapes if sh.name == name)


def test_builder_writes_native_editable_objects_without_pictures(built):
    b, data = built
    prs = Presentation(io.BytesIO(data))
    assert (prs.slide_width, prs.slide_height) == (18288000, 10287000)
    kinds = [it['kind'] for it in b.manifest[0]]
    assert kinds.count('chart') == 3 and kinds.count('table') == 1 and 'picture' not in kinds
    assert not any(sh.shape_type == 13 for sh in prs.slides[0].shapes)  # nenhuma imagem
    title = shape(prs, 'text:title')
    assert title.has_text_frame and title.text_frame.text == 'Título'
    bp = title._element.txBody.find(qn('a:bodyPr'))
    assert bp.get('wrap') == 'none' and bp.get('lIns') == '0' and bp.find(qn('a:noAutofit')) is not None
    rpr = title._element.find('.//' + qn('a:rPr'))
    assert rpr.get('sz') == '5000' and rpr.find(qn('a:latin')).get('typeface') == 'Montserrat Black'
    assert title._element.find('.//' + qn('a:spcPts')).get('val') == str(round(73.34 * 75))
    # linha de base: caixa deslocada para a 1ª linha do PowerPoint (0,75·L) coincidir com a do HTML
    assert title.top == round((60 + 60 - 0.75 * 73.34) * 9525)
    assert 'apps.total' in title._element.find('.//' + qn('p:cNvPr')).get('descr')  # proveniência
    lst = shape(prs, 'list:zone-list')
    p0 = lst._element.find('.//' + qn('a:pPr'))
    assert p0.find(qn('a:buChar')).get('char') == '•' and p0.find(qn('a:buSzPct')).get('val') == '130000'
    assert lst._element.find('.//' + qn('a:br')) is not None  # quebra de linha preservada como quebra (não parágrafo)
    assert shape(prs, 'text:section')._element.txBody.find(qn('a:bodyPr')).get('vert') == 'vert270'
    assert prs.slides[0].notes_slide.notes_text_frame.text == 'Nota sintética do apresentador'


def test_builder_shapes_use_theme_colors_alpha_custgeom_and_texture(built):
    _, data = built
    prs = Presentation(io.BytesIO(data))
    pill = shape(prs, 'rect:pill')._element
    assert pill.find('.//' + qn('a:prstGeom')).get('prst') == 'roundRect'
    assert pill.find('.//' + qn('a:alpha')).get('val') == '18000'
    assert pill.find(qn('p:style')) is None  # sem sombra/estilo implícito do tema do PowerPoint
    card = shape(prs, 'rect:card')._element
    assert card.find('.//' + qn('a:custGeom')) is not None  # raios diferentes por canto
    assert card.find('.//' + qn('a:prstDash')).get('val') == 'dash'
    path = shape(prs, 'path:shape:painel')._element
    assert path.findall('.//' + qn('a:cubicBezTo')) and path.find('.//' + qn('a:close')) is not None
    assert shape(prs, 'rect:unknown')._element.find('.//' + qn('a:pattFill')).get('prst') == 'dkUpDiag'
    assert shape(prs, 'arrow:cmp')._element.find('.//' + qn('a:tailEnd')).get('type') == 'arrow'


def test_builder_matrix_is_native_table_with_full_long_text(built):
    b, data = built
    prs = Presentation(io.BytesIO(data))
    tbl = shape(prs, 'table:matrix')
    assert tbl.has_table
    long_cell = ('Ação sintética ' * 40).strip()
    assert len(long_cell) > 400 and tbl.table.cell(1, 1).text_frame.text == long_cell  # sem truncar
    tc = tbl.table.cell(1, 1)._tc
    # espaço entre células emulado por bordas na cor do fundo, margens = recuo do HTML + meia borda
    assert tc.find('.//' + qn('a:lnL')).get('w') == str(16 * 9525)
    assert tc.find(qn('a:tcPr')).get('marL') == str(round((16 + 8) * 9525))
    assert tbl.table._tbl.tblPr.find(qn('a:tableStyleId')).text == '{2D5ABB26-0587-4C30-8999-92F81FD0307C}'


def test_table_cell_margins_never_force_powerpoint_to_grow_the_row(tmp_path):
    """Cabeçalho da matriz (F07.1): marT 21 + entrelinha 24 + marB 20 > linha de 64 px ⇒ o PowerPoint crescia a linha
    (+5 px medidos) e empurrava a matriz. A margem do lado oposto à âncora é zerada/reduzida sem mover o texto."""
    def cell(valign: str, x: float) -> dict:
        return {'cell': {'x': x, 'y': 300, 'w': 200, 'h': 48}, 'box': {'x': x + 16, 'y': 313, 'w': 168, 'h': 24},
                'paras': [para([run(t=f'Cabeçalho {valign}', size=18.67)], lh=24)], 'wrap': True, 'vertical': False,
                'lines': 1, 'lineWidths': [], 'maxWidth': 168, 'base': 19, 'fill': {'rgb': '4B1A6E', 'a': 1},
                'pad': [13, 16, 12, 16], 'borders': [None, None, None, None], 'valign': valign, 'xid': None}
    op = {'op': 'table', 'key': 'table:matrix', 'box': {'x': 100, 'y': 300, 'w': 648, 'h': 48}, 'kind': 'matrix',
          'cols': [200, 200, 200], 'gap': 16, 'gapColor': {'rgb': 'FFFFFF', 'a': 1}, 'ir': None,
          'rows': [{'y': 300, 'h': 48, 'header': True, 'cells': [cell('t', 100), cell('ctr', 316), cell('b', 532)]}]}
    layout = [{'version': 1, 'index': 1, 'slide_id': 's1', 'canvas': {'width': 1920, 'height': 1080},
               'ops': [{'op': 'background', 'key': 'background', 'fill': None, 'image': None}, op], 'text_nodes': 0}]
    b = pptx_builder.build(pptx_builder.BuildInput(view=synthetic_view(), layouts=layout,
                                                   fontbook=fake_book(tmp_path)))
    tbl = next(sh for sh in b.prs.slides[0].shapes if sh.has_table).table
    shift = 19 - pptx_builder.PPT_BASELINE * 24  # correção de linha de base (igual à das caixas de texto)
    mar = lambda i: {k: int(tbl.cell(0, i)._tc.find(qn('a:tcPr')).get(k)) / 9525 for k in ('marT', 'marB')}  # noqa: E731
    assert mar(0) == pytest.approx({'marT': 13 + 8 + shift, 'marB': 0}, abs=0.01)  # topo: posição intacta
    assert mar(2) == pytest.approx({'marT': 0, 'marB': 12 + 8 - shift}, abs=0.01)  # base: posição intacta
    c = mar(1)  # centro: só a DIFERENÇA entre as margens desloca o texto — preservada
    assert min(c.values()) == pytest.approx(0, abs=0.01)
    assert c['marT'] - c['marB'] == pytest.approx((13 + 8 + shift) - (12 + 8 - shift), abs=0.01)
    assert tbl.rows[0].height == round(64 * 9525)


def test_builder_charts_are_native_with_exact_values_and_empty_gaps(built):
    b, data = built
    prs = Presentation(io.BytesIO(data))
    combo = shape(prs, 'chart:plot')
    assert combo.has_chart
    cs = combo.chart._chartSpace
    assert cs.find('.//' + qn('c:lineChart')) is not None and cs.find('.//' + qn('c:barChart')) is not None
    bar_ax = [a.get('val') for a in cs.find('.//' + qn('c:barChart')).findall(qn('c:axId'))]
    line_ax = [a.get('val') for a in cs.find('.//' + qn('c:lineChart')).findall(qn('c:axId'))]
    assert bar_ax == line_ax  # mesmo eixo (sem eixo secundário)
    assert cs.find('.//' + qn('c:dispBlanksAs')).get('val') == 'gap'
    assert pptx_charts.chart_values(combo) == [[None, 139.0, 140.0], [None, 20.0, 1234.0]]  # lacuna nunca vira 0
    wb = combo.chart.part.chart_workbook.xlsx_part.blob
    assert validate._workbook_values(wb) == [[None, 139.0, 140.0], [None, 20.0, 1234.0]]
    labels = [t.text for t in cs.iter(qn('a:t'))]
    assert '1.234' in labels and 'lacuna' not in labels  # rótulo pt-BR da IR; lacuna sem rótulo
    va = cs.find('.//' + qn('c:valAx'))
    assert va.find('.//' + qn('c:max')).get('val') == '2000.0' and va.find(qn('c:majorUnit')).get('val') == '500.0'
    donut = shape(prs, 'chart:donut')
    dn = donut.chart._chartSpace.find('.//' + qn('c:doughnutChart'))
    assert dn.find(qn('c:holeSize')).get('val') == '60' and dn.find(qn('c:firstSliceAng')).get('val') == '0'
    colors = [p.format.fill.fore_color.rgb for p in donut.chart.plots[0].series[0].points]
    assert str(colors[0]) == 'C00000' and str(colors[1]) == 'F26B1D'
    assert list(donut.chart.plots[0].categories) == ['Crítica', 'Alta', 'Desconhecida']
    sev = shape(prs, 'chart:sev-rows')
    sc = sev.chart._chartSpace
    assert sc.find('.//' + qn('c:barDir')).get('val') == 'bar'
    assert sc.find('.//' + qn('c:catAx')).find('.//' + qn('c:orientation')).get('val') == 'maxMin'
    assert sc.find('.//' + qn('c:valAx')).find(qn('c:delete')).get('val') in (None, '1', 'true')  # eixo oculto
    pts = sc.findall('.//' + qn('c:dPt'))
    assert pts[2].find('.//' + qn('a:pattFill')) is not None  # Desconhecida com textura
    assert pptx_charts.chart_values(sev) == [[2.0, 0.0, 1.0]]  # zero REAL aprovado continua zero


def test_axis_units_follow_html_rules():
    assert pptx_charts.nice_max([139, 140]) == 200 and pptx_charts.nice_max([]) == 1
    assert pptx_charts.axis_units(200, 4, True) == (50, None)
    assert pptx_charts.axis_units(50, 4, True) == (25, 12.5)  # HTML só rotula 0/25/50; grade a cada 12,5
    assert pptx_charts.axis_units(5, 4, True) == (5, 1.25)
    assert pptx_charts.axis_units(2, 4, False) == (0.5, None)


def test_package_is_deterministic_and_clean(built):
    b, data = built
    again = io.BytesIO()
    b.prs.save(again)
    assert opc.normalize(again.getvalue(), pptx_builder._ts(synthetic_view())) == data
    z = zipfile.ZipFile(io.BytesIO(data))
    assert {i.date_time for i in z.infolist()} == {opc.FIXED_ZIP_TIME}
    core = z.read('docProps/core.xml').decode()
    assert '2026-10-16T12:00:00Z' in core
    custom = z.read('docProps/custom.xml').decode()
    assert 'afrika-modelo@2' in custom and 'executive' in custom
    inner = zipfile.ZipFile(io.BytesIO(z.read(next(n for n in z.namelist() if n.endswith('.xlsx')))))
    assert {i.date_time for i in inner.infolist()} == {opc.FIXED_ZIP_TIME}
    assert '2026-10-16T12:00:00Z' in inner.read('docProps/core.xml').decode()
    assert validate.package(data)['macros'] is False


def _zip(entries: dict[str, bytes]) -> bytes:
    buf = io.BytesIO()
    with zipfile.ZipFile(buf, 'w') as z:
        for k, v in entries.items():
            z.writestr(k, v)
    return buf.getvalue()


def test_package_validation_rejects_macros_external_links_and_bad_xml(built):
    _, data = built
    z = zipfile.ZipFile(io.BytesIO(data))
    base = {n: z.read(n) for n in z.namelist()}
    for mutate in (lambda e: e.update({'ppt/vbaProject.bin': b'x'}),
                   lambda e: e.update({'ppt/slides/_rels/x.xml.rels': b'<Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships"><Relationship Id="r" Type="t" Target="http://evil" TargetMode="External"/></Relationships>'}),
                   lambda e: e.update({'ppt/slides/slide1.xml': b'<p:sld'})):
        entries = dict(base)
        mutate(entries)
        with pytest.raises(ExportError) as exc:
            validate.package(_zip(entries))
        assert exc.value.code == 'pptx_invalid'
    with pytest.raises(ExportError):
        validate.package(b'nao e zip')


# ---------------------------------------------------------------- fontes (OFL) e EOT

def test_font_faces_map_css_weights_to_powerpoint_names(tmp_path):
    assets = tmp_path / 'assets'
    assets.mkdir()
    for fam, w in (('montserrat', 900), ('montserrat', 400), ('poppins', 600), ('poppins', 700)):
        (assets / f'{fam}-latin-{w}-normal-abc.woff').write_bytes(b'x')
    (assets / 'montserrat-latin-ext-900-normal-zz.woff').write_bytes(b'x')
    book = F.FontBook(tmp_path)
    assert book.face('Montserrat', 900).typeface == 'Montserrat Black'
    assert (book.face('Poppins', 700).typeface, book.face('Poppins', 700).bold) == ('Poppins', True)
    assert book.face('Poppins', 600).typeface == 'Poppins SemiBold'
    assert (book.face('Segoe UI', 600).typeface, book.face('Segoe UI', 600).bold) == ('Segoe UI', True)
    assert book.file_for_typeface('Montserrat Black', False).name.startswith('montserrat-latin-900')
    assert book.file_for_typeface('Poppins', True).name.startswith('poppins-latin-700')
    assert book.file_for_typeface('Segoe UI', False) is None


def _tiny_sfnt() -> bytes:
    """Fonte TrueType mínima SINTÉTICA (só tabelas lidas pelo EOT) para testar a conversão."""
    name_recs = [(1, 'Fonte Sintetica'), (2, 'Regular'), (4, 'Fonte Sintetica Regular'), (5, 'Version 1.0')]
    strings = b''.join(s.encode('utf-16-be') for _, s in name_recs)
    recs, off = b'', 0
    for nid, s in name_recs:
        b = s.encode('utf-16-be')
        recs += struct.pack('>HHHHHH', 3, 1, 0x409, nid, len(b), off)
        off += len(b)
    name = struct.pack('>HHH', 0, len(name_recs), 6 + 12 * len(name_recs)) + recs + strings
    os2 = bytearray(96)
    os2[0:2] = struct.pack('>H', 4)
    os2[4:6] = struct.pack('>H', 400)
    head = bytearray(54)
    head[12:16] = struct.pack('>I', 0x5F0F3CF5)
    tables = {b'OS/2': bytes(os2), b'head': bytes(head), b'name': name}
    woff_body, entries, offset = b'', b'', 44 + 20 * len(tables)
    import zlib
    for tag in sorted(tables):
        raw = tables[tag]
        comp = zlib.compress(raw)
        entries += struct.pack('>4sIIII', tag, offset + len(woff_body), len(comp), len(raw), 0)
        woff_body += comp + b'\0' * (-len(comp) % 4)
    header = struct.pack('>4sIIHHIHHIIIII', b'wOFF', 0x10000, 44 + len(entries) + len(woff_body), len(tables), 0,
                         0, 1, 0, 0, 0, 0, 0, 0)
    return F.woff_to_sfnt(header + entries + woff_body)


def test_woff_to_ttf_checksum_and_eot_layout():
    sfnt = _tiny_sfnt()
    total = sum(struct.unpack(f'>{len(sfnt) // 4}I', sfnt)) & 0xFFFFFFFF
    assert total == 0xB1B0AFBA  # checkSumAdjustment recalculado
    eot = F.to_eot(sfnt)
    eot_size, data_size, version, flags = struct.unpack('<IIII', eot[:16])
    assert eot_size == len(eot) and data_size == len(sfnt) and version == 0x00020002 and flags == 0
    assert struct.unpack('<H', eot[34:36])[0] == 0x504C and eot.endswith(sfnt)
    size = struct.unpack('<H', eot[82:84])[0]
    assert eot[84:84 + size].decode('utf-16-le') == 'Fonte Sintetica\0'  # família terminada em NUL (como o PowerPoint)
    assert F.info(sfnt)['family'] == 'Fonte Sintetica' and F.embeddable(0) and not F.embeddable(0x2)


# ---------------------------------------------------------------- PDF

def _pdf(pages: int = 2, font_file: bool = True, extra: bytes = b'', box: str = '0 0 1440 810') -> bytes:
    objs = [b'<< /Type /Catalog /Pages 2 0 R >>', f'<< /Type /Pages /Count {pages} >>'.encode()]
    for _ in range(pages):
        objs.append(f'<< /Type /Page /MediaBox [{box}] >>'.encode())
    objs.append(b'<< /Type /FontDescriptor /FontName /AAAAAA+Poppins-Regular '
                + (b'/FontFile2 99 0 R' if font_file else b'') + b' >>')
    fd = len(objs)
    objs.append(f'<< /Type /Font /Subtype /CIDFontType2 /BaseFont /AAAAAA+Poppins-Regular /FontDescriptor {fd} 0 R >>'.encode())
    objs.append(b'<< /Info (sintetico) /CreationDate (D:20200101000000+00\'00\') >>' + extra)
    out, offsets = b'%PDF-1.4\n', []
    for i, o in enumerate(objs, start=1):
        offsets.append(len(out))
        out += f'{i} 0 obj\n'.encode() + o + b'\nendobj\n'
    xref = len(out)
    out += f'xref\n0 {len(objs) + 1}\n0000000000 65535 f \n'.encode()
    out += b''.join(f'{o:010d} 00000 n \n'.encode() for o in offsets)
    out += f'trailer\n<< /Size {len(objs) + 1} >>\nstartxref\n{xref}\n%%EOF\n'.encode()
    return out


def test_pdf_validation_checks_pages_size_fonts_and_active_content():
    canvas = {'width': 1920, 'height': 1080}
    ok = pdfmod.validate(_pdf(), slides=2, canvas=canvas, expected_families={'Poppins'})
    assert ok['pages'] == 2 and ok['page_size_pt'] == [1440.0, 810.0] and ok['fonts'] == ['Poppins']
    for data, slides, fams in ((_pdf(pages=1), 2, None), (_pdf(font_file=False), 2, None),
                               (_pdf(box='0 0 612 792'), 2, None), (_pdf(extra=b' /JavaScript (x)'), 2, None),
                               (_pdf(), 2, {'Montserrat'}), (b'nao e pdf', 2, None)):
        with pytest.raises(ExportError) as exc:
            pdfmod.validate(data, slides=slides, canvas=canvas, expected_families=fams)
        assert exc.value.code == 'pdf_invalid'


def test_pdf_dates_are_normalized_without_moving_offsets():
    from datetime import datetime
    raw = _pdf()
    norm = pdfmod.normalize(raw, datetime(2026, 10, 16, 12, 0, 0))
    assert len(norm) == len(raw) and b"D:20261016120000+00'00'" in norm
    pdfmod.validate(norm, slides=2, canvas={'width': 1920, 'height': 1080}, expected_families=None)


# ---------------------------------------------------------------- documento: matriz > 400 e "Crítica" (12, 10)

def test_matrix_cells_accept_long_text_while_other_lists_keep_their_limit(env):
    c = env.client('Cliente Matriz F07')
    _, p, _ = ready_presentation(env, c['id'])
    doc = env.document(p['id'])
    mx = next(s for s in doc['slides'] if s['layout'] == 'matrix')
    long_cell = 'Plano sintético de remediação com etapas detalhadas. ' * 18
    assert 900 < len(long_cell.strip()) <= 1000
    mx['fields']['items'] = ['Tema | Situação | Ação | Prazo', f'Tema 1 | {long_cell.strip()} | Ação curta | 30 dias']
    env.save(p['id'], doc, 1)  # linha com > 400 caracteres aceita (limite é por célula)
    mx['fields']['items'][1] = f'Tema 1 | {"x" * 1001} | Ação | 30 dias'
    r = env.save(p['id'], doc, 2, status=422)
    assert any(d['type'] == 'item_too_long' for d in r['error']['details'])
    bullets = next(s for s in doc['slides'] if s['fields'].get('items') and s['layout'] != 'matrix')
    mx['fields']['items'][1] = 'Tema 1 | ok | ok | ok'
    bullets['fields']['items'] = ['y' * 401]
    r = env.save(p['id'], doc, 2, status=422)
    assert any(d['type'] == 'item_too_long' and '400' in d['message'] for d in r['error']['details'])


def test_executive_labels_rename_only_very_high_and_keep_the_origin(env):
    c = env.client('Cliente Rótulos F07')
    _, p, v1 = ready_presentation(env, c['id'])
    doc = env.document(p['id'])
    assert 'severity_labels' not in doc  # padrão omitido (hash de documentos antigos inalterado)
    doc['severity_labels'] = 'executive'
    env.save(p['id'], doc, 1)
    v2 = env.view(p['id'])
    assert v1['severity_labels'] == 'veracode' and v2['severity_labels'] == 'executive'

    def severities(v):
        out = []
        for s in v['slides']:
            for z in ((s.get('data') or {}).get('zones') or {}).values():
                if (z.get('data') or {}).get('kind') == 'severity':
                    out.append(z['data'])
        return out

    old, new = severities(v1)[0], severities(v2)[0]
    assert old['categories'][0] == 'Muito alta' and new['categories'][0] == 'Crítica'
    assert new['categories'][1:] == old['categories'][1:]  # "Muito baixa" etc. preservadas
    p0 = new['series'][0]['points'][0]
    assert p0['source_label'] == 'Muito alta' and p0['key'] == 'VERY_HIGH' and p0['y'] == old['series'][0]['points'][0]['y']
    assert p0['color'] == 'severity.VERY_HIGH'  # cor continua a do nível (vermelho)
    assert charts.severity_label('VERY_LOW', 'executive') == 'Muito baixa'


def test_export_plan_matches_f07_object_choices(env):
    c = env.client('Cliente Plano F07')
    _, p, v = ready_presentation(env, c['id'])
    plan = export_plan.plan(v)
    assert plan['status'] == 'IMPLEMENTED'
    assert export_plan.MATRIX['ranking_bars']['export_as'] == 'table'
    empty = {'kind': 'bar', 'available': True, 'empty': True, 'series': []}
    assert export_plan._chart('bar_chart', 'bar', empty, 'data')[0]['component'] == 'unavailable'
    assert all(s['data_as_image'] == [] for s in plan['slides'])


def test_trend_colors_require_declared_direction_not_the_math_sign():
    up = charts.trend(SAST, 5, None)  # backlog aberto (menor é melhor): +5 = piora comprovada
    down = charts.trend(SAST, -3, None)
    neutral = charts.trend('apps.total', 5, None)  # direção não declarada => neutra, apesar do sinal positivo
    assert (up['trend'], up['trend_color']) == ('worsening', 'variation.worsening')
    assert (down['trend'], down['trend_color']) == ('improvement', 'variation.improvement')
    assert (neutral['trend'], neutral['trend_color']) == ('unknown', 'variation.neutral')


def test_full_values_cover_every_afrika_v2_metric():
    from app.presentations import templates
    from app.presentations.document import TemplateDefinition, all_metrics
    defn = TemplateDefinition.model_validate(next(d for (k, v), d in templates.BUILTIN.items()
                                                  if k == 'afrika-modelo' and v == 2))
    assert {m for s in defn.slides for m in all_metrics(s)} <= set(full_values())


def test_system_font_stacks_use_the_theme_export_family(tmp_path):
    v = synthetic_view()
    v['theme']['fonts'] = {r: {'stack': ['system-ui', 'Segoe UI', 'sans-serif'], 'export_family': 'Arial'}
                           for r in ('heading', 'body', 'numeric')}
    lay = synthetic_layout()
    lay[0]['ops'] = [op for op in lay[0]['ops'] if op['op'] in ('background', 'text')]
    for op in lay[0]['ops']:
        for p in op.get('paras', []):
            for r in p['runs']:
                r['font'] = 'system-ui'
    b = pptx_builder.build(pptx_builder.BuildInput(view=v, layouts=lay, fontbook=fake_book(tmp_path)))
    faces = {el.get('typeface') for el in b.prs.slides[0].shapes._spTree.iter(qn('a:latin'))}
    assert faces == {'Arial'}  # genérico CSS nunca vai para o PPTX: usa o export_family declarado pelo tema
