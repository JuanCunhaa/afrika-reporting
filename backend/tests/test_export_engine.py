"""F07 — exportação REAL (Chromium headless + build do frontend): PPTX/PDF gerados, reabertos e inspecionados.

Sem o motor (navegador ou build do renderizador), os testes são PULADOS com o motivo; com
`AFRIKA_EXPORT_TESTS=required` (job de exportação do CI) eles FALHAM em vez de pular — nada fica oculto.
Dados 100% SINTÉTICOS; nenhuma Veracode; ativos de marca sintéticos em pasta descartável.
"""
from __future__ import annotations

import io
import json
import os
import struct
import sys
import time
import uuid
import zipfile
import zlib
from dataclasses import dataclass

import pytest
from fastapi.testclient import TestClient
from pptx import Presentation
from pptx.oxml.ns import qn

from app import db
from app.export import browser, renderer, service as export_service, validate
from app.export.errors import ExportError
from app.main import create_app
from conftest import sqlite_url
from export_fixtures import SAST, ready_presentation
from test_presentations import OP, Env


def _engine_status() -> tuple[bool, str]:
    try:
        renderer.locate()
        browser.discover()
    except ExportError as exc:
        return False, f'{exc.code}: {exc.message}'
    return True, ''


ENGINE_OK, ENGINE_REASON = _engine_status()
REQUIRED = os.getenv('AFRIKA_EXPORT_TESTS') == 'required'
pytestmark = pytest.mark.skipif(not ENGINE_OK and not REQUIRED,
                                reason=f'motor de exportação indisponível ({ENGINE_REASON})')
AFRIKA_V2 = {'VERY_HIGH': 'C00000', 'HIGH': 'F26B1D', 'MEDIUM': 'FFC000', 'LOW': '2F6FD6', 'VERY_LOW': '2F6FD6',
             'INFORMATIONAL': '8C8C8C', 'UNKNOWN': 'B5B5B5'}


@dataclass
class World:
    api: TestClient
    env: Env
    client: dict
    sub: dict
    p_client: dict
    p_sub: dict
    brand: object


def _png(w: int, h: int, rgb: tuple[int, int, int]) -> bytes:
    def chunk(t: bytes, d: bytes) -> bytes:
        return struct.pack('>I', len(d)) + t + d + struct.pack('>I', zlib.crc32(t + d) & 0xFFFFFFFF)
    raw = b''.join(b'\0' + bytes(rgb) * w for _ in range(h))
    return (b'\x89PNG\r\n\x1a\n' + chunk(b'IHDR', struct.pack('>IIBBBBB', w, h, 8, 2, 0, 0, 0))
            + chunk(b'IDAT', zlib.compress(raw)) + chunk(b'IEND', b''))


@pytest.fixture(scope='module')
def world(tmp_path_factory):
    if not ENGINE_OK:
        pytest.fail(f'AFRIKA_EXPORT_TESTS=required, mas o motor está indisponível: {ENGINE_REASON}')
    tmp = tmp_path_factory.mktemp('export-world')
    url = sqlite_url(tmp / 'db.sqlite3')
    db.upgrade(url)
    brand = tmp / 'brand'
    (brand / 'afrika').mkdir(parents=True)
    mp = pytest.MonkeyPatch()
    mp.setenv('AFRIKA_BRAND_ASSET_DIR', str(brand))
    mp.setenv('AFRIKA_EXPORT_DIR', str(tmp / 'exports'))
    try:
        with TestClient(create_app(url)) as api:
            env = Env(api)
            env.builtin()
            client = env.client('Cliente Exportação Sintético')
            sub = env.subclient(client['id'], 'Unidade Sintética Leste')
            _, p_client, _ = ready_presentation(env, client['id'])
            _, p_sub, _ = ready_presentation(env, sub['id'], kind='subclient')
            yield World(api, env, client, sub, p_client, p_sub, brand)
    finally:
        mp.undo()


def export(api: TestClient, pid: str, fmt: str, revision: int = 1, timeout: float = 120) -> tuple[dict, bytes]:
    r = api.post(f'/api/presentations/{pid}/exports', json={'format': fmt, 'revision': revision, 'operator': OP})
    assert r.status_code == 202, r.text
    job = r.json()
    end = time.monotonic() + timeout
    while job['status'] in ('queued', 'running') and time.monotonic() < end:
        time.sleep(0.1)
        job = api.get(f'/api/exports/{job["id"]}').json()
    if job['status'] != 'done':
        return job, b''
    f = api.get(job['download_url'])
    assert f.status_code == 200
    return job, f.content


def by_name(slide) -> dict:
    out: dict = {}
    for sh in slide.shapes:
        out.setdefault(sh.name, sh)
    return out


@pytest.fixture(scope='module')
def afrika(world):
    job, data = export(world.api, world.p_client['id'], 'pptx')
    assert job['status'] == 'done', json.dumps(job.get('error'), ensure_ascii=False)
    view = world.env.view(world.p_client['id'])
    return job, data, view


# ---------------------------------------------------------------- PPTX afrika-modelo@2 (1, 3–9, 14)

def test_pptx_is_valid_editable_and_follows_the_document(afrika):
    job, data, view = afrika
    rep = job['report']
    assert rep['package']['macros'] is False and rep['package']['external_relationships'] is False
    assert rep['content']['plan_match'] is True and rep['content']['images'] == 0
    assert rep['content']['objects']['chart'] >= 6 and rep['content']['objects']['table'] >= 4
    assert data[:2] == b'PK' and zipfile.ZipFile(io.BytesIO(data)).testzip() is None
    prs = Presentation(io.BytesIO(data))
    assert (prs.slide_width, prs.slide_height) == (18288000, 10287000)  # 16:9, 1920×1080 px
    assert len(prs.slides) == len(view['slides']) == 16
    for slide, vs in zip(prs.slides, view['slides']):  # mesma ordem e textos editoriais
        shapes = by_name(slide)
        title = shapes.get('text:title')
        assert title is not None and title.has_text_frame, vs['slide_id']
        assert title.text_frame.text.replace('\v', ' ').replace('\n', ' ') == vs['fields']['title'], vs['slide_id']
    assert not any(sh.shape_type == 13 for s in prs.slides for sh in s.shapes)  # nenhuma imagem (sem ativos)


def test_kpis_are_text_and_charts_are_native_with_embedded_exact_values(afrika):
    _, data, view = afrika
    prs = Presentation(io.BytesIO(data))
    kpi = next(sh for s in prs.slides for sh in s.shapes if sh.name == 'text:kpi-value:apps.total')
    assert kpi.has_text_frame and kpi.text_frame.text == '140'
    descr = kpi._element.find('.//' + qn('p:cNvPr')).get('descr')
    assert 'apps.total' in descr and 'revisão aprovada' in descr  # proveniência no texto alternativo
    charts_ = [(i, sh) for i, s in enumerate(prs.slides) for sh in s.shapes if getattr(sh, 'has_chart', False)]
    assert charts_
    gaps = 0
    for i, sh in charts_:
        vs = view['slides'][i]
        assert sh.chart.part.chart_workbook.xlsx_part is not None  # dados incorporados (planilha)
        for zone in (vs.get('data') or {}).get('zones', {}).values():
            ir = zone.get('data') or {}
            if ir.get('kind') in ('combo', 'line') and ir.get('available') is not False:
                expected = [[p['y'] if p['state'] == 'VALUE' else None for p in s['points']] for s in ir['series']]
                got = [list(ser.values) for plot in sh.chart.plots for ser in plot.series]
                if got == expected:
                    gaps += sum(v is None for s in got for v in s)
                    assert list(sh.chart.plots[0].categories) == ir['categories']
    assert gaps > 0  # meses sem aprovação continuam VAZIOS (nunca 0)


def test_tables_are_native_with_the_approved_ranking(afrika):
    _, data, _ = afrika
    prs = Presentation(io.BytesIO(data))
    tables = [sh for s in prs.slides for sh in s.shapes if sh.has_table]
    texts = {tbl.table.cell(r, c).text_frame.text for tbl in tables for r in range(len(tbl.table.rows))
             for c in range(len(tbl.table.columns))}
    assert 'Aplicação Sintética 01' in texts and '64' in texts
    assert 'VULN-SINT-001' in texts


def test_severity_colors_follow_afrika_v2_semantics(afrika):
    _, data, view = afrika
    prs = Presentation(io.BytesIO(data))
    donut = next(sh for s in prs.slides for sh in s.shapes if sh.name.startswith('chart:') and getattr(sh, 'has_chart', False)
                 and sh.chart._chartSpace.find('.//' + qn('c:doughnutChart')) is not None)
    colors = [str(p.format.fill.fore_color.rgb) for p in donut.chart.plots[0].series[0].points]
    assert colors == list(AFRIKA_V2.values())  # Crítica vermelho … Baixa e Muito baixa azul, Informativa cinza
    assert list(donut.chart.plots[0].categories)[:2] == ['Muito alta', 'Alta']  # rótulos Veracode por padrão


def test_variation_pills_use_proven_trend_colors(afrika):
    _, data, _ = afrika
    prs = Presentation(io.BytesIO(data))
    shapes = {sh.name: sh for s in prs.slides for sh in s.shapes}

    def dot(metric: str) -> str:
        sh = next(v for k, v in shapes.items() if k.startswith(f'rect:cmp-delta:{metric}#2')
                  or k.startswith(f'rect:pill:{metric}#2'))
        return str(sh.fill.fore_color.rgb)

    def text(metric: str) -> str:
        return next(v.text_frame.text for k, v in shapes.items()
                    if k.startswith((f'text:cmp-delta:{metric}', f'text:pill:{metric}')))

    assert dot(SAST) == 'C00000' and 'piora' in text(SAST)  # backlog aberto subiu: piora comprovada
    assert dot('apps.total') == '7A7A7A'  # direção não declarada: neutra, mesmo com variação positiva


def test_export_is_deterministic_for_the_same_revision(world, afrika):
    job, data, _ = afrika
    again, data2 = export(world.api, world.p_client['id'], 'pptx')
    assert again['sha256'] == job['sha256'] and data2 == data
    pdf1, b1 = export(world.api, world.p_client['id'], 'pdf')
    pdf2, b2 = export(world.api, world.p_client['id'], 'pdf')
    assert pdf1['status'] == pdf2['status'] == 'done' and b1 == b2 and pdf1['sha256'] == pdf2['sha256']


def test_pdf_is_valid_with_theme_fonts_embedded(world):
    job, data = export(world.api, world.p_client['id'], 'pdf')
    assert job['status'] == 'done', json.dumps(job.get('error'), ensure_ascii=False)
    rep = job['report']['pdf']
    assert data.startswith(b'%PDF-1.') and data.rstrip().endswith(b'%%EOF')
    assert rep['pages'] == 16 and rep['page_size_pt'] == [1440.0, 810.0] and rep['fonts_embedded'] is True
    assert {f.lower()[:6] for f in rep['fonts']} <= {'montse', 'poppin', 'beviet'}
    assert job['filename'].endswith('_r1.pdf') and job['timings']['total_s'] < 60


def test_custom_properties_record_template_theme_and_sources(afrika):
    _, data, view = afrika
    custom = zipfile.ZipFile(io.BytesIO(data)).read('docProps/custom.xml').decode()
    for expected in ('afrika-modelo@2', 'afrika@2', view['revision']['revision_hash'],
                     view['sources']['primary']['content_hash'][:16], 'veracode'):
        assert expected in custom


# ---------------------------------------------------------------- cliente × subcliente, revisões (16, 17, 20)

def test_subclient_export_is_isolated_from_the_client(world, afrika):
    _, client_data, _ = afrika
    job, data = export(world.api, world.p_sub['id'], 'pptx')
    assert job['status'] == 'done'
    prs = Presentation(io.BytesIO(data))
    assert 'Unidade Sintética Leste' in prs.core_properties.subject
    assert f'subclient:{world.sub["id"]}' in zipfile.ZipFile(io.BytesIO(data)).read('docProps/custom.xml').decode()
    client_prs = Presentation(io.BytesIO(client_data))
    assert 'Unidade Sintética Leste' not in client_prs.core_properties.subject


def test_each_persisted_revision_exports_its_own_content(world):
    c = world.env.client('Cliente Revisões Exportação')
    _, p, _ = ready_presentation(world.env, c['id'], template='demo-executivo', version=1)
    doc = world.env.document(p['id'])
    doc['title'] = 'Título salvo na revisão 2'
    world.env.save(p['id'], doc, 1)
    j1, d1 = export(world.api, p['id'], 'pptx', revision=1)
    j2, d2 = export(world.api, p['id'], 'pptx', revision=2)
    assert j1['status'] == j2['status'] == 'done' and j1['revision'] == 1 and j2['revision'] == 2
    assert Presentation(io.BytesIO(d1)).core_properties.title != 'Título salvo na revisão 2'
    assert Presentation(io.BytesIO(d2)).core_properties.title == 'Título salvo na revisão 2'


# ---------------------------------------------------------------- templates antigos (15)

@pytest.mark.parametrize('template, version', [('demo-executivo', 1), ('afrika-modelo', 1)])
def test_previous_templates_still_export(world, template, version):
    c = world.env.client(f'Cliente {template} v{version}')
    _, p, view = ready_presentation(world.env, c['id'], template=template, version=version)
    for fmt in ('pptx', 'pdf'):
        job, data = export(world.api, p['id'], fmt)
        assert job['status'] == 'done', (fmt, json.dumps(job.get('error'), ensure_ascii=False))
        assert data
    assert len(Presentation(io.BytesIO(export(world.api, p['id'], 'pptx')[1])).slides) == len(view['slides'])


# ---------------------------------------------------------------- matriz > 400, overflow, rótulos (10, 12, 13)

def _matrix_presentation(world, name: str, cell: str, rows: int):
    c = world.env.client(name)
    _, p, _ = ready_presentation(world.env, c['id'])
    doc = world.env.document(p['id'])
    mx = next(s for s in doc['slides'] if s['layout'] == 'matrix')
    mx['fields']['items'] = ['Tema | Situação atual | Ação recomendada | Próximo passo'] + [
        f'Tema {i + 1} | {cell} | {cell} | {cell}' for i in range(rows)]
    world.env.save(p['id'], doc, 1)
    return p, mx


def test_matrix_with_long_cells_exports_full_text_with_reduced_font(world):
    cell = ('Situação sintética detalhada com contexto, impacto e responsáveis definidos; acompanhamento '
            'quinzenal com evidências anexadas e critérios objetivos de encerramento. ') * 3
    cell = cell.strip()
    assert len(cell) > 400
    p, mx = _matrix_presentation(world, 'Cliente Matriz Longa', cell, 2)
    job, data = export(world.api, p['id'], 'pptx', revision=2)
    assert job['status'] == 'done', json.dumps(job.get('error'), ensure_ascii=False)
    prs = Presentation(io.BytesIO(data))
    tbl = next(sh for s in prs.slides for sh in s.shapes if sh.has_table and sh.name.startswith('table:matrix'))
    assert tbl.table.cell(1, 1).text_frame.text == cell  # conteúdo integral, nunca truncado
    sizes = {int(r.get('sz')) for r in tbl.table.cell(1, 1)._tc.iter(qn('a:rPr'))}
    assert max(sizes) < round(18.67 * 75)  # fonte reduzida (corpo 14 pt) para caber
    assert min(sizes) >= round(18.67 * 75 * 0.6) - 1  # nunca abaixo do mínimo (60%)


def test_content_that_cannot_fit_blocks_export_with_slide_diagnostic(world):
    p, _ = _matrix_presentation(world, 'Cliente Matriz Excessiva', 'z' * 6 + (' texto extenso sintético' * 40), 8)
    job, data = export(world.api, p['id'], 'pptx', revision=2)
    assert job['status'] == 'failed' and data == b''
    assert job['error']['code'] == 'content_overflow'
    assert any('matriz' in d['field'] and d['type'] == 'clipped-y' for d in job['error']['details'])


def _bu_values(first_name: str) -> dict:
    names = [first_name, 'Unidade Sul', 'Unidade Leste', 'Unidade Oeste', 'Unidade Centro']
    return {'apps.by_business_unit': {'items': [{'business_unit_id': f'ab{k}', 'name': n, 'applications': 41 - k * 7}
                                                for k, n in enumerate(names)], 'unknown': 0,
                                      'distinct_applications': 135}}


def _bu_plot_height(data: bytes, first: str) -> float:
    for s in Presentation(io.BytesIO(data)).slides:
        for sh in s.shapes:
            if sh.has_chart and first in list(sh.chart.plots[0].categories):
                ml = sh.chart._chartSpace.chart.plotArea.find(qn('c:layout') + '/' + qn('c:manualLayout'))
                return float(ml.find(qn('c:h')).get('val'))
    raise AssertionError(f'gráfico com a categoria {first!r} ausente')


def test_long_axis_category_wraps_in_full_in_html_and_pptx(world, afrika):
    """F07.1: categoria longa QUEBRA em linhas no HTML/PDF (antes virava "Unidade Sintét…") e chega inteira ao PPTX."""
    name = 'Unidade Sintética Norte Ampliada'
    c = world.env.client('Cliente Rótulo Longo')
    _, p, _ = ready_presentation(world.env, c['id'], values=_bu_values(name))
    job, data = export(world.api, p['id'], 'pptx')
    assert job['status'] == 'done', json.dumps(job.get('error'), ensure_ascii=False)  # abreviação bloquearia
    # mesma área do HTML, que reservou as linhas extras do rótulo: área de plotagem menor que a de nomes curtos
    assert _bu_plot_height(data, name) < _bu_plot_height(afrika[1], 'Unidade Norte') - 0.02
    job, data = export(world.api, p['id'], 'pdf')
    assert job['status'] == 'done', json.dumps(job.get('error'), ensure_ascii=False)


def test_axis_category_that_cannot_fit_blocks_export_instead_of_abbreviating(world):
    c = world.env.client('Cliente Rótulo Excessivo')
    _, p, _ = ready_presentation(world.env, c['id'], values=_bu_values('Unidade sintética ' * 12))
    job, data = export(world.api, p['id'], 'pdf')
    assert job['status'] == 'failed' and data == b''
    assert job['error']['code'] == 'content_overflow'
    assert any(d['type'] == 'ellipsis' and 'rótulo do eixo' in d['message'] for d in job['error']['details'])


def test_executive_labels_reach_the_pptx_with_origin(world):
    c = world.env.client('Cliente Rótulo Executivo')
    _, p, _ = ready_presentation(world.env, c['id'])
    doc = world.env.document(p['id'])
    doc['severity_labels'] = 'executive'
    world.env.save(p['id'], doc, 1)
    job, data = export(world.api, p['id'], 'pptx', revision=2)
    assert job['status'] == 'done'
    prs = Presentation(io.BytesIO(data))
    cats = [list(sh.chart.plots[0].categories) for s in prs.slides for sh in s.shapes
            if getattr(sh, 'has_chart', False) and sh.chart._chartSpace.find('.//' + qn('c:doughnutChart')) is not None]
    assert cats and all(c[0] == 'Crítica' and 'Muito baixa' in c for c in cats)
    assert 'executive' in zipfile.ZipFile(io.BytesIO(data)).read('docProps/custom.xml').decode()


# ---------------------------------------------------------------- fontes ausentes, ativos, erros (22, 23, 24)

def test_missing_theme_fonts_block_the_export(world):
    with world.api.app.state.session_factory() as s:
        from app.presentations import service as ps
        pres = ps.get(s, uuid.UUID(world.p_client['id']))
        v = export_service.precheck(s, pres, ps.get_revision(s, pres, 1))
    with pytest.raises(ExportError) as exc:
        export_service.build_export(v, 'pdf', block_fonts=True)  # simula arquivos de fonte ausentes no build
    assert exc.value.code == 'font_unavailable'
    assert exc.value.details and all(d['type'] == 'font_fallback' for d in exc.value.details)


def test_brand_assets_become_the_only_pictures_and_absence_falls_back(world):
    c = world.env.client('Cliente Ativos Sintéticos')
    _, p, _ = ready_presentation(world.env, c['id'])
    job, data = export(world.api, p['id'], 'pptx')
    assert job['status'] == 'done' and not any(
        sh.shape_type == 13 for s in Presentation(io.BytesIO(data)).slides for sh in s.shapes)  # sem ativo: fallback
    (world.brand / 'afrika' / 'logo-conteudo.png').write_bytes(_png(104, 68, (93, 36, 130)))
    (world.brand / 'afrika' / 'logo-capa.svg').write_text(
        '<svg xmlns="http://www.w3.org/2000/svg" width="644" height="126"><rect width="644" height="126" '
        'fill="#5D2482"/></svg>', encoding='utf-8')
    try:
        job, data = export(world.api, p['id'], 'pptx')
        assert job['status'] == 'done', json.dumps(job.get('error'), ensure_ascii=False)
        prs = Presentation(io.BytesIO(data))
        pics = [(i, sh) for i, s in enumerate(prs.slides) for sh in s.shapes if sh.shape_type == 13]
        assert pics and {sh.name for _, sh in pics} == {'brand:logo-conteudo', 'brand:logo-capa'}
        assert all('Ativo visual autorizado' in sh._element.find('.//' + qn('p:cNvPr')).get('descr') for _, sh in pics)
        svg = next(sh for _, sh in pics if sh.name == 'brand:logo-capa')
        assert svg._element.find('.//{http://schemas.microsoft.com/office/drawing/2016/SVG/main}svgBlip') is not None
        assert job['report']['content']['images'] == len(pics)
    finally:
        for f in (world.brand / 'afrika').iterdir():
            f.unlink()


def test_unusable_browser_fails_the_job_explicitly(world, monkeypatch):
    monkeypatch.setenv(browser.BROWSER_ENV, sys.executable)  # executável existente que não é um navegador
    job, data = export(world.api, world.p_client['id'], 'pdf', timeout=60)
    assert job['status'] == 'failed' and data == b'' and job['download_url'] is None
    assert job['error']['code'] in ('export_engine_failed', 'export_timeout')


def test_pptx_passes_independent_revalidation(afrika):
    job, data, view = afrika
    assert validate.package(data)['embedded_workbooks'] == job['report']['content']['objects']['chart']
