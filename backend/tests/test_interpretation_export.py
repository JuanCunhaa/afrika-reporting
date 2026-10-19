"""F08 — texto de interpretação APROVADO e aplicado ao slide sai igual no HTML (visualização), no PPTX e no PDF, com
o motor de exportação REAL e sem nenhuma chamada a provedor durante a exportação. Dados 100% sintéticos.

Mesma política dos testes de exportação (F07): sem motor => SKIP local; `AFRIKA_EXPORT_TESTS=required` (CI) => FALHA.
"""
from __future__ import annotations

import io
import re
import zlib

import pytest
from pptx import Presentation as PptxFile

from app.interpretation import providers
from export_fixtures import ready_presentation
from test_export_engine import ENGINE_OK, ENGINE_REASON, REQUIRED, export
from test_interpretation import approve, gen
from test_presentations import OP, Env

pytestmark = pytest.mark.skipif(not ENGINE_OK and not REQUIRED,
                                reason=f'motor de exportação indisponível ({ENGINE_REASON})')


def _squash(s: str) -> str:
    return re.sub(r'\s+', '', s)


def _cmaps(data: bytes) -> dict[int, dict[int, str]]:
    """Fonte (número do objeto) -> mapa glifo->texto pelo /ToUnicode (Type0/Identity-H do Chromium)."""
    objs = {int(n): body for n, body in re.findall(rb'(\d+) 0 obj(.*?)endobj', data, re.S)}

    def stream(body: bytes) -> bytes:
        raw = body.split(b'stream', 1)[1].lstrip(b'\r\n').rsplit(b'endstream', 1)[0]
        return zlib.decompress(raw) if b'/FlateDecode' in body.split(b'stream', 1)[0] else raw

    out = {}
    for n, body in objs.items():
        m = re.search(rb'/ToUnicode (\d+) 0 R', body)
        if not m or int(m.group(1)) not in objs:
            continue
        cmap = stream(objs[int(m.group(1))])
        table: dict[int, str] = {}
        for block in re.findall(rb'beginbfchar(.*?)endbfchar', cmap, re.S):
            for src, dst in re.findall(rb'<([0-9A-Fa-f]+)>\s*<([0-9A-Fa-f]+)>', block):
                table[int(src, 16)] = bytes.fromhex(dst.decode()).decode('utf-16-be')
        for block in re.findall(rb'beginbfrange(.*?)endbfrange', cmap, re.S):
            for a, b, dst in re.findall(rb'<([0-9A-Fa-f]+)>\s*<([0-9A-Fa-f]+)>\s*<([0-9A-Fa-f]+)>', block):
                base = int(dst, 16)
                for i, code in enumerate(range(int(a, 16), int(b, 16) + 1)):
                    table[code] = chr(base + i)
        out[n] = table
    return out


def pdf_text(data: bytes) -> str:
    """Texto do PDF (sem espaços) decodificado pelas CMaps — só para conferir o conteúdo; stdlib apenas."""
    cmaps = _cmaps(data)
    names = {}
    for block in re.findall(rb'/Font\s*<<(.*?)>>', data, re.S):
        for name, num in re.findall(rb'/(\w+) (\d+) 0 R', block):
            names[name.decode()] = int(num)
    text = []
    for body in re.findall(rb'obj\s*<<(.*?)>>\s*stream\r?\n(.*?)endstream', data, re.S):
        hdr, raw = body
        try:
            content = zlib.decompress(raw) if b'/FlateDecode' in hdr else raw
        except zlib.error:
            continue
        if b'Tf' not in content:
            continue
        table: dict[int, str] = {}
        for tok in re.finditer(rb'/(\w+) [\d.]+ Tf|<([0-9A-Fa-f]+)>', content):
            if tok.group(1):
                table = cmaps.get(names.get(tok.group(1).decode(), -1), {})
            elif table:
                hx = tok.group(2).decode()
                text += [table.get(int(hx[i:i + 4], 16), '') for i in range(0, len(hx), 4)]
    return _squash(''.join(text))


def pptx_text(data: bytes, index: int) -> str:
    slide = PptxFile(io.BytesIO(data)).slides[index]
    return _squash(''.join(sh.text_frame.text for sh in slide.shapes if sh.has_text_frame))


@pytest.fixture(scope='module')
def applied(tmp_path_factory):
    from app import db
    from app.main import create_app
    from conftest import sqlite_url
    from fastapi.testclient import TestClient
    if not ENGINE_OK:
        pytest.fail(f'AFRIKA_EXPORT_TESTS=required, mas o motor está indisponível: {ENGINE_REASON}')
    tmp = tmp_path_factory.mktemp('interp-export')
    url = sqlite_url(tmp / 'db.sqlite3')
    db.upgrade(url)
    mp = pytest.MonkeyPatch()
    mp.setenv('AFRIKA_EXPORT_DIR', str(tmp / 'exports'))
    try:
        with TestClient(create_app(url)) as api:
            env = Env(api)
            env.builtin()
            c = env.client('Cliente Interpretação Exportada')
            _, p, view = ready_presentation(env, c['id'], previous=('2026-07', '2026-08'))
            pid = p['id']
            body = approve(api, pid, gen(api, pid, 'painel-kpis', 'EXECUTIVE_SUMMARY')['id'])
            r = api.post(f'/api/presentations/{pid}/interpretations/{body["id"]}/apply',
                         json={'operator': OP, 'expected_revision': 1, 'overwrite_confirmed': True})
            assert r.status_code == 200, r.text
            items = approve(api, pid, gen(api, pid, 'painel-graficos', 'ATTENTION_POINTS', 'TECHNICAL')['id'])
            r = api.post(f'/api/presentations/{pid}/interpretations/{items["id"]}/apply',
                         json={'operator': OP, 'expected_revision': 2, 'overwrite_confirmed': True})
            assert r.status_code == 200, r.text
            # Exportação NUNCA chama provedor: qualquer chamada a partir daqui derruba o teste.
            def forbidden(*a, **k):
                raise AssertionError('provedor de interpretação chamado durante a exportação')
            mp.setattr(providers.DeterministicProvider, 'generate', forbidden)
            mp.setattr(providers.ExternalProvider, 'generate', forbidden)
            yield api, env, pid, body, items, view
    finally:
        mp.undo()


def _index(view: dict, slide_id: str) -> int:
    return [s['slide_id'] for s in view['slides']].index(slide_id)


def test_html_pptx_and_pdf_show_the_same_persisted_interpretation(applied):
    api, env, pid, body, items, _ = applied
    view = env.view(pid)
    assert view['presentation']['current_revision'] == 3 and view['ready'] is True
    html = {s['slide_id']: s['fields'] for s in view['slides']}
    assert html['painel-kpis']['body'] == body['text']
    assert html['painel-graficos']['items'] == [s['text'] for s in items['statements']]

    job, pptx = export(api, pid, 'pptx', revision=3)
    assert job['status'] == 'done', job.get('error')  # cabe no slide: nenhum bloqueio de corte/overflow
    assert _squash(body['text']) in pptx_text(pptx, _index(view, 'painel-kpis'))
    graf = pptx_text(pptx, _index(view, 'painel-graficos'))
    assert all(_squash(s['text']) in graf for s in items['statements'])

    job, pdf = export(api, pid, 'pdf', revision=3)
    assert job['status'] == 'done', job.get('error')
    text = pdf_text(pdf)
    assert _squash(body['text']) in text
    assert all(_squash(s['text']) in text for s in items['statements'])


def test_previous_revision_exports_without_the_interpretation(applied):
    api, env, pid, body, _, _ = applied
    job, pptx = export(api, pid, 'pptx', revision=1)
    assert job['status'] == 'done', job.get('error')
    assert _squash(body['text']) not in pptx_text(pptx, _index(env.view(pid), 'painel-kpis'))
    job, pdf = export(api, pid, 'pdf', revision=1)
    assert job['status'] == 'done' and _squash(body['text']) not in pdf_text(pdf)
    assert 'Indicadoresprincipais' in pdf_text(pdf) or len(pdf_text(pdf)) > 500  # extrator funciona no mesmo PDF
