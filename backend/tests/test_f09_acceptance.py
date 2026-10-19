"""F09 — homologação ponta a ponta com dados 100% SINTÉTICOS (tenants fixture `gamma` e `beta`; nenhuma Veracode real).

Um único fluxo, na ordem da ordem do PO (etapas 1–22), cada etapa num teste próprio para a evidência ficar por etapa:
cadastro → conexões herdada/independente → escopos → seleção de aplicações/usuários → coleta simulada → falhas
(indisponibilidade/incompletude) → competência → snapshot → valor manual justificado → aprovação → histórico →
template Afrika → apresentação → edição → interpretação local → revisão humana → aplicação → PPTX/PDF → download →
REINÍCIO do backend sobre o MESMO arquivo SQLite (nova instância da aplicação) sem perda do que foi persistido.

As etapas de exportação seguem a política F07: sem motor => SKIP local; `AFRIKA_EXPORT_TESTS=required` => FALHA.
Os testes dependem da ordem (estado compartilhado no módulo); rodar o arquivo inteiro.
"""
from __future__ import annotations

import getpass
import hashlib
import io
import json
import logging
import re
import sqlite3
import time
import zipfile
from contextlib import closing
from dataclasses import dataclass, field
from pathlib import Path
from typing import Any

import httpx
import pytest
from fastapi.testclient import TestClient
from pptx import Presentation as PptxFile

from app import db
from app.interpretation import providers
from app.main import create_app
from app.veracode.fixtures import FixtureVeracode, build_tenant, synthetic_credentials
from app.veracode.settings import VeracodeSettings
from conftest import SleepRecorder, error_code, sqlite_url, veracode_environ
from test_export_engine import ENGINE_OK, ENGINE_REASON, REQUIRED
from test_interpretation_export import pdf_text, pptx_text
from test_metrics import gamma_tenant, ids, sev

OP = 'Operadora Homologação Sintética'
APPROVER = 'Aprovadora Homologação Sintética'
REVIEWER = 'Revisora Homologação Sintética'
JUST = 'Valor sintético declarado na homologação F09 (fonte indisponível no tenant fixture)'
LABEL = '2026-09'

needs_engine = pytest.mark.skipif(not ENGINE_OK and not REQUIRED,
                                  reason=f'motor de exportação indisponível ({ENGINE_REASON})')


@dataclass
class World:
    url: str
    tmp: Any
    tenant: Any
    fake: FixtureVeracode
    faults: dict
    mp: pytest.MonkeyPatch
    ctx: TestClient | None = None
    api: TestClient | None = None
    s: dict = field(default_factory=dict)  # estado compartilhado entre as etapas

    def start(self) -> TestClient:
        def handler(request: httpx.Request) -> httpx.Response:
            if self.faults.get('sca_500') and request.url.path.endswith('/findings') \
                    and request.url.params.get('scan_type') == 'SCA':
                return httpx.Response(500, headers={'Content-Type': 'application/json'})
            return self.fake.handle(request)

        app = create_app(self.url, veracode_settings=VeracodeSettings(page_size=50, max_retries=2),
                         veracode_transport=httpx.MockTransport(handler),
                         environ=veracode_environ(VC_GAMMA='gamma', VC_BETA='beta'), sleep=SleepRecorder())
        self.ctx = TestClient(app)
        self.api = self.ctx.__enter__()
        return self.api

    def stop(self) -> None:
        if self.ctx is not None:
            self.ctx.__exit__(None, None, None)
        self.ctx = self.api = None

    # ---- atalhos
    def call(self, method: str, url: str, expected: int = 200, **body) -> dict:
        r = self.api.request(method, url, json=body or None)
        assert r.status_code == expected, f'{method} {url} -> {r.status_code}: {r.text}'
        return r.json()

    def lock(self, cid: str, n: int) -> int:
        return self.call('GET', f'/api/competences/{cid}/revisions/{n}')['lock_version']

    def run(self, kind: str, owner_id: str) -> dict:
        r = self.call('POST', '/api/metrics/runs', 202, owner_kind=kind, owner_id=owner_id, competence=LABEL)
        end = time.monotonic() + 30
        while time.monotonic() < end:
            body = self.call('GET', f'/api/metrics/runs/{r["run_id"]}')
            if body['status'] in ('done', 'cancelled', 'failed'):
                return body
            time.sleep(0.02)
        raise AssertionError('coleta não terminou')

    def approve_manual(self, owner_id: str, label: str, values: dict) -> str:
        cid = self.call('POST', '/api/competences', 201, owner_kind='client', owner_id=owner_id, competence=label,
                        operator=OP)['id']
        base = f'/api/competences/{cid}/revisions/1'
        for mid, v in values.items():
            self.call('PUT', f'{base}/metrics/{mid}', value=v, justification=JUST, operator=OP,
                      expected_lock_version=self.lock(cid, 1))
        self.call('PUT', f'{base}/required', metric_ids=['apps.total'], operator=OP,
                  expected_lock_version=self.lock(cid, 1))
        self.call('POST', f'{base}/submit', operator=OP, expected_lock_version=self.lock(cid, 1))
        self.call('POST', f'{base}/approve', approver=APPROVER, expected_lock_version=self.lock(cid, 1))
        return cid

    def interp(self, pid: str, iid: str, action: str, expected: int = 200, **body) -> dict:
        return self.call('POST', f'/api/presentations/{pid}/interpretations/{iid}/{action}', expected,
                         operator=REVIEWER, **body)

    def export(self, pid: str, fmt: str, revision: int) -> tuple[dict, httpx.Response]:
        job = self.call('POST', f'/api/presentations/{pid}/exports', 202, format=fmt, revision=revision, operator=OP)
        end = time.monotonic() + 180
        while job['status'] in ('queued', 'running') and time.monotonic() < end:
            time.sleep(0.1)
            job = self.call('GET', f'/api/exports/{job["id"]}')
        assert job['status'] == 'done', json.dumps(job.get('error'), ensure_ascii=False)
        f = self.api.get(job['download_url'])
        assert f.status_code == 200
        return job, f


def M(run: dict, metric_id: str) -> dict:
    return next(m for m in run['result']['metrics'] if m['metric_id'] == metric_id)


def snap(detail: dict, metric_id: str) -> dict:
    return next(s for s in detail['snapshots'] if s['metric_id'] == metric_id)


@pytest.fixture(scope='module')
def w(tmp_path_factory):
    tmp = tmp_path_factory.mktemp('f09')
    url = sqlite_url(tmp / 'f09.sqlite3')
    db.upgrade(url)
    tenant = gamma_tenant()
    fake = FixtureVeracode({'gamma': tenant,
                            'beta': build_tenant('beta', apps=3, bus=1, teams=1, users=1, label='Beta')})
    mp = pytest.MonkeyPatch()
    mp.setenv('AFRIKA_EXPORT_DIR', str(tmp / 'exports'))
    world = World(url, tmp, tenant, fake, {}, mp)
    # todos os logs do fluxo inteiro ficam guardados para a etapa de segurança (23)
    world.s['logs'] = logs = []
    handler = logging.Handler(logging.DEBUG)
    handler.emit = lambda record: logs.append(record.getMessage())
    root = logging.getLogger()
    previous = root.level
    root.addHandler(handler)
    root.setLevel(logging.DEBUG)
    world.start()
    try:
        yield world
    finally:
        world.stop()
        root.removeHandler(handler)
        root.setLevel(previous)
        mp.undo()


# ---------------------------------------------------------------- 1–3. cadastro e conexões

def test_01_client_registration(w):
    gamma = w.call('POST', '/api/connections', 201, name='Conexão Sintética Gama', region='commercial',
                   credential_ref='env:VC_GAMMA')
    beta = w.call('POST', '/api/connections', 201, name='Conexão Sintética Beta', region='commercial',
                  credential_ref='env:VC_BETA')
    c = w.call('POST', '/api/clients', 201, name='Cliente Homologação F09', connection_id=gamma['id'],
               timezone='UTC')
    solo = w.call('POST', '/api/clients', 201, name='Cliente Isolado F09', connection_id=beta['id'], timezone='UTC')
    assert c['connection_id'] == gamma['id'] and c['timezone'] == 'UTC'
    # credencial nunca trafega: só a referência env:NOME
    assert 'api_key' not in json.dumps(gamma).lower() and gamma['credential_ref'] == 'env:VC_GAMMA'
    w.s.update(gamma=gamma, beta=beta, client=c, solo=solo)


def test_02_subclient_registration(w):
    cid = w.s['client']['id']
    inherited = w.call('POST', f'/api/clients/{cid}/subclients', 201, name='Unidade Sintética Herdada')
    own = w.call('POST', f'/api/clients/{cid}/subclients', 201, name='Unidade Sintética Independente',
                 connection_id=w.s['beta']['id'])
    subs = w.call('GET', f'/api/clients/{cid}/subclients')['items']
    assert {s['id'] for s in subs} == {inherited['id'], own['id']}
    # cliente sem subclientes continua válido
    assert w.call('GET', f'/api/clients/{w.s["solo"]["id"]}/subclients')['items'] == []
    w.s.update(sub=inherited, own=own)


def test_03_inherited_and_independent_connection(w):
    inh = w.call('GET', f'/api/veracode/subclients/{w.s["sub"]["id"]}')
    own = w.call('GET', f'/api/veracode/subclients/{w.s["own"]["id"]}')
    assert inh['scope']['connection_source'] == 'client' and inh['scope']['connection']['id'] == w.s['gamma']['id']
    assert own['scope']['connection_source'] == 'subclient' and own['scope']['connection']['id'] == w.s['beta']['id']
    assert inh['credential_status'] == own['credential_status'] == 'configured'
    inh_t = w.call('POST', f'/api/veracode/subclients/{w.s["sub"]["id"]}/test')
    own_t = w.call('POST', f'/api/veracode/subclients/{w.s["own"]["id"]}/test')
    assert inh_t['status'] == own_t['status'] == 'ok'
    # valores de credencial nunca aparecem nas respostas (só a referência env:NOME)
    for tenant in ('gamma', 'beta'):
        for secret in synthetic_credentials(tenant):
            assert secret not in json.dumps([inh, own, inh_t, own_t])
    gamma_apps = {a['guid'] for a in w.call('GET', f'/api/veracode/subclients/{w.s["sub"]["id"]}/applications'
                                                 '?limit=200')['items']}
    beta_apps = {a['guid'] for a in w.call('GET', f'/api/veracode/subclients/{w.s["own"]["id"]}/applications'
                                                '?limit=200')['items']}
    assert gamma_apps == set(ids(w.tenant)['apps']) and beta_apps and not gamma_apps & beta_apps


# ---------------------------------------------------------------- 4–5. escopos e seleção

def test_04_scopes(w):
    cid, sid = w.s['client']['id'], w.s['sub']['id']
    I = ids(w.tenant)
    w.call('PUT', f'/api/scopes/clients/{cid}/applications', 201,
           rule={'field': 'business_unit.id', 'cmp': 'in', 'values': I['bu']}, expected_version=0)
    w.call('PUT', f'/api/scopes/subclients/{sid}/applications', 201,
           rule={'field': 'business_unit.id', 'cmp': 'eq', 'value': I['bu'][0]}, expected_version=0)
    w.call('PUT', f'/api/scopes/clients/{cid}/users', 201, rule={'op': 'all'}, expected_version=0)
    w.call('PUT', f'/api/scopes/subclients/{sid}/users', 201,
           rule={'field': 'team.id', 'cmp': 'in', 'values': [I['teams'][0]]}, expected_version=0)
    # domínios não se misturam: campo de aplicação recusado em usuários
    r = w.api.put(f'/api/scopes/subclients/{sid}/users',
                  json={'rule': {'field': 'business_unit.id', 'cmp': 'eq', 'value': I['bu'][0]},
                        'expected_version': 1})
    assert r.status_code == 422 and error_code(r) == 'invalid_rule'
    # concorrência otimista: versão esperada desatualizada não sobrescreve
    r = w.api.put(f'/api/scopes/subclients/{sid}/applications',
                  json={'rule': {'op': 'all'}, 'expected_version': 0})
    assert r.status_code == 409
    assert w.call('GET', f'/api/scopes/clients/{cid}')  # versões ativas consultáveis


def test_05_application_and_user_selection(w):
    cid, sid = w.s['client']['id'], w.s['sub']['id']
    g = ids(w.tenant)['apps']
    apps = w.call('POST', f'/api/scope-preview/clients/{cid}', domain='applications', sample_limit=200)
    sub = next(s for s in apps['subclients'] if s['id'] == sid)
    assert {r['id'] for r in sub['sample']} == {g[0], g[2]}  # BU0 = App 001 e 003
    assert apps['universe']['included'] == 4
    users = w.call('POST', f'/api/scope-preview/clients/{cid}', domain='users', sample_limit=200)
    team0 = ids(w.tenant)['teams'][0]
    members = {m['user_id'] for m in w.tenant.team_users[team0]}
    active = {u['user_id'] for u in w.tenant.users if u['active']}
    sub_u = next(s for s in users['subclients'] if s['id'] == sid)
    assert {r['id'] for r in sub_u['sample']} == members & active


# ---------------------------------------------------------------- 6–7. coleta simulada e falhas

def test_06_simulated_collection(w):
    run_c = w.run('client', w.s['client']['id'])
    run_s = w.run('subclient', w.s['sub']['id'])
    assert run_c['status'] == run_s['status'] == 'done'
    assert M(run_c, 'apps.total')['status'] == 'VALIDATED' and M(run_c, 'apps.total')['value'] == 4
    assert M(run_s, 'apps.total')['value'] == 2
    # valores-ouro do tenant gamma (test_metrics.py): Policy SAST abertos = 5 no cliente; 3 no subcliente (App 001)
    assert M(run_c, 'sast.policy.open.by_severity')['value'] == sev(VERY_HIGH=1, HIGH=2, MEDIUM=1, INFORMATIONAL=1)
    assert M(run_s, 'sast.policy.open.by_severity')['value']['total'] == 3
    prov = run_c['result']['provenance']  # proveniência da execução: versões de escopo/contrato e hash da entrada
    assert prov['run_id'] == run_c['run_id'] and prov['input_hash'] and prov['scope_refs']
    for m in run_c['result']['metrics']:  # todo indicador traz status, filtros e temporalidade
        assert m['status'] and m['temporal'] is not None and m['filters'] is not None
        assert m['publishable'] == (m['status'] == 'VALIDATED')
    w.s.update(run_c=run_c, run_s=run_s)


def test_07_unavailability_and_incompleteness(w):
    run_c = w.s['run_c']
    # identidade não comprovada => INCOMPLETE sem valor definitivo; fechamento histórico => UNAVAILABLE
    distinct = M(run_c, 'sast.open.distinct')
    assert distinct['status'] == 'INCOMPLETE' and distinct['value'] is None
    assert all(m['value'] is None for m in run_c['result']['metrics'] if m['status'] != 'VALIDATED')
    assert M(run_c, 'sast.policy.open.at_period_end')['status'] == 'UNAVAILABLE'
    # subcliente sem escopo próprio (conexão independente) => UNAVAILABLE, nunca zero
    run_own = w.run('subclient', w.s['own']['id'])
    total = M(run_own, 'apps.total')
    assert total['status'] == 'UNAVAILABLE' and total['value'] is None and total['reasons']
    # HTTP 500 persistente na Findings SCA => indicador SCA sem valor (não vira 0); SAST segue validado
    w.call('DELETE', f'/api/veracode/clients/{w.s["client"]["id"]}/cache')  # sem cache: a coleta vai à fonte
    w.faults['sca_500'] = True
    try:
        faulty = w.run('client', w.s['client']['id'])
    finally:
        w.faults['sca_500'] = False
        w.call('DELETE', f'/api/veracode/clients/{w.s["client"]["id"]}/cache')  # respostas com falha não ficam
    sca = M(faulty, 'sca.policy.open.by_severity')
    assert sca['status'] in ('INCOMPLETE', 'FAILED') and sca['value'] is None
    assert M(faulty, 'sast.policy.open.by_severity')['status'] == 'VALIDATED'


# ---------------------------------------------------------------- 8–12. competência, snapshot, manual, aprovação, histórico

def test_08_competence_creation(w):
    out = w.call('POST', '/api/competences/from-run', 201, run_id=w.s['run_c']['run_id'], operator=OP)
    cid, n = out['competence_id'], out['revision']
    d = w.call('GET', f'/api/competences/{cid}/revisions/{n}')
    assert d['status'] == 'DRAFT' and d['competence']['competence'] == LABEL
    # importar a coleta de OUTRO dono na mesma competência é recusado
    r = w.api.post(f'/api/competences/{cid}/revisions/{n}/import-run',
                   json={'run_id': w.s['run_s']['run_id'], 'operator': OP,
                         'expected_lock_version': d['lock_version']})
    assert r.status_code == 409 and error_code(r) == 'run_owner_mismatch'
    w.s.update(cid=cid, n=n)


def test_09_persistent_snapshot(w):
    cid, n = w.s['cid'], w.s['n']
    d = w.call('GET', f'/api/competences/{cid}/revisions/{n}')
    total = snap(d, 'apps.total')
    assert total['origin'] == 'VERACODE' and total['value_state'] == 'VALIDATED' and total['value'] == 4
    assert snap(d, 'sast.open.distinct')['value'] is None and snap(d, 'sast.open.distinct')['value_state'] == \
        'INCOMPLETE'
    # persistido no SQLite (não só em memória): leitura direta do arquivo
    path = w.url.removeprefix('sqlite:///')
    with closing(sqlite3.connect(path)) as con:  # Uuid do SQLAlchemy no SQLite: 32 hex sem hífens
        rows = con.execute('SELECT COUNT(*) FROM metric_snapshots ms JOIN competence_revisions r '
                           'ON r.id = ms.revision_id WHERE r.competence_id = ?', (cid.replace('-', ''),)).fetchone()
    assert rows[0] == len(d['snapshots']) > 0


def test_10_justified_manual_fill(w):
    cid, n = w.s['cid'], w.s['n']
    base = f'/api/competences/{cid}/revisions/{n}'
    r = w.api.put(f'{base}/metrics/sast.open.distinct',
                  json={'value': sev(HIGH=1), 'justification': '', 'operator': OP,
                        'expected_lock_version': w.lock(cid, n)})
    assert r.status_code == 422  # justificativa obrigatória
    w.call('PUT', f'{base}/metrics/sast.open.distinct', value=sev(VERY_HIGH=2, HIGH=2, MEDIUM=2, LOW=1,
                                                                    INFORMATIONAL=1),
           justification=JUST, operator=OP, expected_lock_version=w.lock(cid, n))
    w.call('PUT', f'{base}/metrics/licenses.usage', value=0, justification=JUST, operator=OP,
           expected_lock_version=w.lock(cid, n))
    # bloqueio da F03 (subcliente independente sem escopo) deixa `apps.unclassified` INCOMPLETE; o operador declara
    w.call('PUT', f'{base}/metrics/apps.unclassified', value=2, justification=JUST, operator=OP,
           expected_lock_version=w.lock(cid, n))
    d = w.call('GET', base)
    manual = snap(d, 'sast.open.distinct')
    assert manual['origin'] == 'MANUAL' and manual['value']['total'] == 8
    assert manual['source']['status'] == 'INCOMPLETE'  # valor da fonte preservado ao lado do manual
    zero = snap(d, 'licenses.usage')
    assert zero['value'] == 0 and zero['value_state'] != 'MISSING'  # zero é valor, não ausência


def test_11_competence_approval(w):
    cid, n = w.s['cid'], w.s['n']
    base = f'/api/competences/{cid}/revisions/{n}'
    w.call('PUT', f'{base}/required', metric_ids=['apps.total', 'sast.policy.open.by_severity'], operator=OP,
           expected_lock_version=w.lock(cid, n))
    w.call('POST', f'{base}/submit', operator=OP, expected_lock_version=w.lock(cid, n))
    stale = w.lock(cid, n) - 1
    r = w.api.post(f'{base}/approve', json={'approver': APPROVER, 'expected_lock_version': stale})
    assert r.status_code == 409  # lock otimista
    w.call('POST', f'{base}/approve', approver=APPROVER, note='homologação F09', expected_lock_version=w.lock(cid, n))
    d = w.call('GET', base)
    assert d['status'] == 'APPROVED' and not d['editable']
    # versão aprovada é imutável pela API ...
    r = w.api.put(f'{base}/metrics/apps.total', json={'value': 99, 'justification': JUST, 'operator': OP,
                                                      'expected_lock_version': d['lock_version']})
    assert r.status_code == 409
    # ... e pelo banco (gatilhos append-only)
    path = w.url.removeprefix('sqlite:///')
    con = sqlite3.connect(path)
    try:
        with pytest.raises(sqlite3.IntegrityError, match='snapshot_revision_not_draft'):
            con.execute('UPDATE metric_snapshots SET value_json = ? WHERE 1 = 1', ('99',))
            con.commit()
        with pytest.raises(sqlite3.IntegrityError):
            con.execute('DELETE FROM metric_snapshots')
            con.commit()
    finally:
        con.close()
    integ = w.call('GET', f'/api/competences/{cid}/integrity')
    assert integ['ok'] is True, integ
    w.s['approved_detail'] = d


def test_12_approved_history_retrieval(w):
    cid_client = w.s['client']['id']
    w.s['prev'] = {lbl: w.approve_manual(cid_client, lbl, {'apps.total': v, 'sast.policy.open.by_severity':
                                                           sev(HIGH=v - 1, MEDIUM=2)})
                   for lbl, v in (('2026-07', 3), ('2026-08', 5))}
    s = w.call('GET', f'/api/history/client/{cid_client}/series?metric_id=apps.total&months=6')
    got = {p['competence']: p['value'] for p in s['points'] if p['state'] == 'APPROVED'}
    assert got == {'2026-07': 3, '2026-08': 5, '2026-09': 4}
    assert s['gaps'] == ['2026-04', '2026-05', '2026-06'] and s['interpolation'] == 'none'
    cmp_ = w.call('GET', f'/api/history/client/{cid_client}/compare?competence={LABEL}&metric_id=apps.total')
    item = next(i for i in cmp_['items'] if i['metric_id'] == 'apps.total')
    assert cmp_['previous_competence'] == '2026-08' and item['delta'] == -1  # 4 (2026-09) − 5 (2026-08)
    # isolamento: o histórico do cliente isolado não enxerga nada do cliente homologado
    other = w.call('GET', f'/api/history/client/{w.s["solo"]["id"]}/series?metric_id=apps.total&months=6')
    assert all(p['state'] == 'GAP' for p in other['points'])
    assert w.call('GET', f'/api/competences?owner_kind=client&owner_id={w.s["solo"]["id"]}')['items'] == []


# ---------------------------------------------------------------- 13–15. template, apresentação, edição

def test_13_afrika_template_selection(w):
    w.call('POST', '/api/presentation-templates/builtin', operator=OP)
    cid_client = w.s['client']['id']
    sel = w.call('PUT', f'/api/template-selection/client/{cid_client}', template_key='afrika-modelo',
                 template_version=2, operator=OP)
    got = w.call('GET', f'/api/template-selection/client/{cid_client}')
    assert got['template']['key'] == 'afrika-modelo' and got['template']['version'] == 2, (sel, got)
    assert got['selected_by'] == OP
    # o subcliente herda a escolha do cliente; o cliente isolado não
    sub = w.call('GET', f'/api/template-selection/subclient/{w.s["sub"]["id"]}')
    assert sub['template']['key'] == 'afrika-modelo' and sub['inherited'] is True
    solo = w.call('GET', f'/api/template-selection/client/{w.s["solo"]["id"]}')
    assert solo['template'] is None or solo['template']['key'] != 'afrika-modelo'


def test_14_presentation_creation(w):
    cid = w.s['cid']
    av = w.call('GET', f'/api/presentation-templates/afrika-modelo/versions/2/availability?competence_id={cid}')
    # Escolha do operador: slides completos/editoriais + parciais sem indicador obrigatório (lacuna = aviso).
    # Parcial COM indicador obrigatório bloquearia a exportação (seria preciso nova revisão aprovada).
    slots = [s['slot_id'] for s in av['slots']
             if s['required'] or s['status'] in ('complete', 'editorial')
             or (s['status'] == 'partial' and not s['required_metrics'])]
    assert 'painel-kpis' in slots, [(s['slot_id'], s['status'], s['required_metrics'],
                                     [m for m in s['metrics'] if m['state'] != 'VALUE']) for s in av['slots']]
    w.s['availability'] = {s['slot_id']: s['status'] for s in av['slots']}
    p = w.call('POST', '/api/presentations', 201, competence_id=cid, operator=OP, slots=slots)
    assert p['template']['key'] == 'afrika-modelo' and p['template']['version'] == 2, p
    view = w.call('GET', f'/api/presentations/{p["id"]}/view')
    assert view['ready'] is True and view['blockers'] == 0, [i for s in view['slides'] for i in s['issues']]
    # lacunas preservadas: indicadores sem valor aparecem como aviso/lacuna, nunca como número
    issues = [i for s in view['slides'] for i in s['issues']]
    assert any(i['code'] == 'manual_value' for i in issues)
    assert any(i['code'] == 'history_gap' for i in issues)  # 2026-04..06 sem aprovação: lacuna, sem interpolação
    w.s.update(pid=p['id'], view1=view, slots=slots)


def test_15_slide_editing(w):
    pid = w.s['pid']
    rev = w.call('GET', f'/api/presentations/{pid}')['current_revision']
    doc = w.call('GET', f'/api/presentations/{pid}/revisions/{rev}')['document']
    target = next(s for s in doc['slides'] if s['template_slot'] == 'capa' or s['slide_id'].startswith('capa'))
    target['fields']['title'] = 'Relatório Sintético de Homologação'
    out = w.call('PUT', f'/api/presentations/{pid}/document', document=doc, expected_revision=rev, operator=OP,
                 summary='homologação F09: título da capa')
    assert out['current_revision'] == rev + 1
    r = w.api.put(f'/api/presentations/{pid}/document',
                  json={'document': doc, 'expected_revision': rev, 'operator': OP, 'summary': 'conflito'})
    assert r.status_code == 409  # edição concorrente não sobrescreve
    old = w.call('GET', f'/api/presentations/{pid}/revisions/{rev}')['document']
    assert next(s for s in old['slides'] if s['slide_id'] == target['slide_id'])['fields'].get('title') != \
        'Relatório Sintético de Homologação'
    w.s['capa'] = target['slide_id']


# ---------------------------------------------------------------- 16–18. interpretação

def test_16_local_interpretation_generation(w):
    pid = w.s['pid']
    cfg = w.call('GET', '/api/interpretation/config')
    # provedor externo de IA desabilitado por padrão; só o provedor local determinístico
    assert cfg['settings']['external_enabled'] is False and cfg['settings']['provider'] == 'deterministic-demo'
    assert cfg['settings']['credential_configured'] is False
    ex = w.call('POST', f'/api/presentations/{pid}/interpretations', 201, slide_id='painel-kpis', level='EXECUTIVE',
                kind='EXECUTIVE_SUMMARY', operator=OP)
    # mesmo slide = mesma cadeia de versões; o texto técnico vai para outro slide
    tech = w.call('POST', f'/api/presentations/{pid}/interpretations', 201, slide_id='painel-graficos',
                  level='TECHNICAL', kind='ATTENTION_POINTS', operator=OP)
    for iv in (ex, tech):
        assert iv['provider'] == 'deterministic-demo' and iv['status'] == 'GENERATED'
        assert iv['validation']['valid'] is True and iv['validation']['blockers'] == 0
    assert ex['text'] != tech['text']
    w.s.update(ex=ex, tech=tech)


def test_17_validation_and_human_review(w):
    pid = w.s['pid']
    # rejeição exige motivo; rejeitada não pode ser aplicada
    w.interp(pid, w.s['tech']['id'], 'review')
    r = w.api.post(f'/api/presentations/{pid}/interpretations/{w.s["tech"]["id"]}/reject',
                   json={'operator': REVIEWER})
    assert r.status_code == 422
    rej = w.interp(pid, w.s['tech']['id'], 'reject', reason='Tom técnico demais para a diretoria (homologação)')
    assert rej['status'] == 'REJECTED'
    # regenerar = nova versão encadeada
    regen = w.call('POST', f'/api/presentations/{pid}/interpretations/{w.s["ex"]["id"]}/regenerate', 201,
                   operator=OP)
    assert regen['id'] != w.s['ex']['id'] and regen['version'] > w.s['ex']['version']
    # só a versão mais recente avança
    r = w.api.post(f'/api/presentations/{pid}/interpretations/{w.s["ex"]["id"]}/review', json={'operator': REVIEWER})
    assert r.status_code == 409
    # número inventado numa edição humana é BLOQUEADO pela validação determinística
    bad = w.call('POST', f'/api/presentations/{pid}/interpretations/{regen["id"]}/edit', 201, operator=REVIEWER,
                 text='O backlog caiu para 9999 achados graças à nova equipe.', justification='teste de bloqueio')
    assert bad['status'] == 'IN_REVIEW'  # edição humana já nasce em revisão
    assert bad['validation']['valid'] is False and bad['validation']['blockers'] >= 1
    assert any(i['code'].startswith('unverified') for i in bad['validation']['issues']), bad['validation']
    r = w.api.post(f'/api/presentations/{pid}/interpretations/{bad["id"]}/approve', json={'operator': REVIEWER})
    assert r.status_code == 409 and error_code(r) == 'interpretation_not_valid'
    # versão nova a partir do texto válido do provedor local
    good = w.call('POST', f'/api/presentations/{pid}/interpretations/{bad["id"]}/regenerate', 201, operator=OP)
    w.interp(pid, good['id'], 'review')
    approved = w.interp(pid, good['id'], 'approve')
    assert approved['status'] == 'APPROVED'
    w.s['approved_iv'] = approved


def test_18_apply_text_to_slide_and_stale_block(w):
    pid, iv = w.s['pid'], w.s['approved_iv']
    rev = w.call('GET', f'/api/presentations/{pid}')['current_revision']
    out = w.interp(pid, iv['id'], 'apply', expected_revision=rev, overwrite_confirmed=True)
    view = w.call('GET', f'/api/presentations/{pid}/view')
    assert view['presentation']['current_revision'] == rev + 1
    assert next(s for s in view['slides'] if s['slide_id'] == 'painel-kpis')['fields']['body'] == iv['text'], out
    # nova revisão APROVADA da competência + rebind => interpretação pendente fica desatualizada e é bloqueada
    pending = w.call('POST', f'/api/presentations/{pid}/interpretations', 201, slide_id='painel-kpis',
                     level='EXECUTIVE', kind='EXECUTIVE_SUMMARY', operator=OP)
    w.interp(pid, pending['id'], 'review')
    cid = w.s['cid']
    rev2 = w.call('POST', f'/api/competences/{cid}/revisions', 201, operator=OP)['number']
    base = f'/api/competences/{cid}/revisions/{rev2}'
    w.call('PUT', f'{base}/metrics/apps.total', value=5, justification=JUST + ' (revisão auditada)', operator=OP,
           expected_lock_version=w.lock(cid, rev2))
    w.call('POST', f'{base}/submit', operator=OP, expected_lock_version=w.lock(cid, rev2))
    w.call('POST', f'{base}/approve', approver=APPROVER, expected_lock_version=w.lock(cid, rev2))
    # revisão 1 aprovada continua íntegra e consultável (sem sobrescrita silenciosa)
    r1 = w.call('GET', f'/api/competences/{cid}/revisions/{w.s["n"]}')
    assert r1['status'] in ('APPROVED', 'SUPERSEDED') and snap(r1, 'apps.total')['value'] == 4
    assert w.call('GET', f'/api/competences/{cid}/official')['number'] == rev2
    cur = w.call('GET', f'/api/presentations/{pid}')['current_revision']
    w.call('POST', f'/api/presentations/{pid}/rebind', expected_revision=cur, operator=OP, revision=rev2)
    r = w.api.post(f'/api/presentations/{pid}/interpretations/{pending["id"]}/approve', json={'operator': REVIEWER})
    assert r.status_code == 409 and error_code(r) == 'interpretation_stale', r.text
    audit = w.call('GET', f'/api/competences/{cid}/audit')['items']
    kinds = [e['event_type'] if 'event_type' in e else e.get('type') for e in audit]
    assert kinds.count('revision.approved') == 2, kinds  # duas aprovações auditadas, nenhuma sobrescrita
    assert w.call('GET', f'/api/competences/{cid}/integrity')['audit_chain']['ok'] is True
    w.s['final_revision'] = w.call('GET', f'/api/presentations/{pid}')['current_revision']
    w.s['final_view'] = w.call('GET', f'/api/presentations/{pid}/view')


def test_18b_isolation_between_clients(w):
    """Cliente isolado (outra conexão, sem subclientes, template neutro) não enxerga nada do cliente homologado."""
    from test_presentations import values_for
    solo = w.s['solo']['id']
    cid = w.approve_manual(solo, LABEL, values_for(0))
    p = w.call('POST', '/api/presentations', 201, competence_id=cid, operator=OP, template_key='demo-executivo')
    assert w.call('GET', f'/api/presentations/{p["id"]}/view')['ready'] is True
    mine = w.call('GET', f'/api/presentations?owner_kind=client&owner_id={w.s["client"]["id"]}')['items']
    theirs = w.call('GET', f'/api/presentations?owner_kind=client&owner_id={solo}')['items']
    assert {x['id'] for x in mine} == {w.s['pid']} and {x['id'] for x in theirs} == {p['id']}
    # interpretação de um cliente não é acessível pelo caminho da apresentação do outro
    iv = w.s['approved_iv']['id']
    assert w.api.get(f'/api/presentations/{p["id"]}/interpretations/{iv}').status_code == 404
    assert w.call('GET', f'/api/presentations/{p["id"]}/interpretations')['items'] == []
    # a apresentação do cliente isolado não usa nenhum valor do cliente homologado
    view = json.dumps(w.call('GET', f'/api/presentations/{p["id"]}/view'), ensure_ascii=False)
    assert 'Cliente Homologação F09' not in view and 'Unidade Sintética' not in view
    # a competência do cliente homologado não é importável/visível como do outro dono
    assert w.call('GET', f'/api/competences?owner_kind=client&owner_id={solo}')['items'][0]['id'] == cid
    w.s.update(solo_pid=p['id'])


# ---------------------------------------------------------------- 19–21. exportação e download

def _pptx_inventory(data: bytes) -> dict:
    deck = PptxFile(io.BytesIO(data))
    out = {'slides': len(deck.slides), 'charts': 0, 'tables': 0, 'pictures': 0, 'text': 0}
    for slide in deck.slides:
        for sh in slide.shapes:
            out['charts'] += sh.has_chart
            out['tables'] += sh.has_table
            out['pictures'] += sh.shape_type == 13
            out['text'] += sh.has_text_frame
    with zipfile.ZipFile(io.BytesIO(data)) as z:
        out['embedded_xlsx'] = sum(1 for n in z.namelist() if n.endswith('.xlsx'))
        assert z.testzip() is None
    return out


@needs_engine
def test_19_pptx_export(w):
    if not ENGINE_OK:
        pytest.fail(f'AFRIKA_EXPORT_TESTS=required, mas o motor está indisponível: {ENGINE_REASON}')
    pid, rev, view = w.s['pid'], w.s['final_revision'], w.s['final_view']
    job, f = w.export(pid, 'pptx', rev)
    inv = _pptx_inventory(f.content)
    assert inv['slides'] == len(view['slides'])
    assert inv['charts'] >= 1 and inv['charts'] == inv['embedded_xlsx'] and inv['pictures'] == 0
    order = [s['slide_id'] for s in view['slides']]
    idx = order.index('painel-kpis')
    assert ''.join(w.s['approved_iv']['text'].split()) in pptx_text(f.content, idx)
    assert ''.join('Relatório Sintético de Homologação'.split()) in pptx_text(f.content, order.index(w.s['capa']))
    w.s.update(pptx_job=job, pptx=f, pptx_inv=inv)


@needs_engine
def test_20_pdf_export(w):
    pid, rev, view = w.s['pid'], w.s['final_revision'], w.s['final_view']
    job, f = w.export(pid, 'pdf', rev)
    assert f.content.startswith(b'%PDF-') and b'%%EOF' in f.content[-1024:]
    pages = len(re.findall(rb'/Type\s*/Page(?![a-zA-Z])', f.content))
    assert pages == len(view['slides']), (pages, len(view['slides']))  # um slide por página, mesma ordem do HTML
    assert ''.join(w.s['approved_iv']['text'].split()) in pdf_text(f.content)
    w.s.update(pdf_job=job, pdf=f)
    # modelo neutro (cliente isolado) também exporta nos dois formatos
    for fmt in ('pptx', 'pdf'):
        job, f = w.export(w.s['solo_pid'], fmt, 1)
        assert job['template'].startswith('demo-executivo') and f.content[:4] in (b'PK\x03\x04', b'%PDF')


@needs_engine
def test_21_download_both_formats(w):
    for key, ext, media in (('pptx', '.pptx', 'presentationml'), ('pdf', '.pdf', 'application/pdf')):
        f, job = w.s[key], w.s[f'{key}_job']
        assert hashlib.sha256(f.content).hexdigest() == job['sha256'] == f.headers['x-export-sha256']
        assert media in f.headers['content-type'] and f.headers['x-content-type-options'] == 'nosniff'
        assert f.headers['cache-control'] == 'no-store' and ext in f.headers['content-disposition']
        assert len(f.content) == job['size']
    # id inexistente não vaza caminho do servidor
    r = w.api.get('/api/exports/00000000-0000-4000-8000-000000000000/file')
    assert r.status_code == 404 and 'tmp' not in r.text.lower() and '\\' not in r.text


# ---------------------------------------------------------------- 22. reabertura do sistema

def test_22_restart_without_losing_persisted_data(w):
    before = {
        'clients': w.call('GET', '/api/clients?limit=200')['items'],
        'subs': w.call('GET', f'/api/clients/{w.s["client"]["id"]}/subclients')['items'],
        'scopes': w.call('GET', f'/api/scopes/clients/{w.s["client"]["id"]}'),
        'comp': w.call('GET', f'/api/competences/{w.s["cid"]}'),
        'pres': w.call('GET', f'/api/presentations/{w.s["pid"]}'),
        'view': w.call('GET', f'/api/presentations/{w.s["pid"]}/view'),
        'interps': w.call('GET', f'/api/presentations/{w.s["pid"]}/interpretations'),
        'history': w.call('GET', f'/api/history/client/{w.s["client"]["id"]}/series?metric_id=apps.total&months=6'),
        'audit': w.call('GET', f'/api/competences/{w.s["cid"]}/audit'),
    }
    w.stop()
    w.start()  # NOVA instância da aplicação sobre o mesmo arquivo SQLite
    after = {
        'clients': w.call('GET', '/api/clients?limit=200')['items'],
        'subs': w.call('GET', f'/api/clients/{w.s["client"]["id"]}/subclients')['items'],
        'scopes': w.call('GET', f'/api/scopes/clients/{w.s["client"]["id"]}'),
        'comp': w.call('GET', f'/api/competences/{w.s["cid"]}'),
        'pres': w.call('GET', f'/api/presentations/{w.s["pid"]}'),
        'view': w.call('GET', f'/api/presentations/{w.s["pid"]}/view'),
        'interps': w.call('GET', f'/api/presentations/{w.s["pid"]}/interpretations'),
        'history': w.call('GET', f'/api/history/client/{w.s["client"]["id"]}/series?metric_id=apps.total&months=6'),
        'audit': w.call('GET', f'/api/competences/{w.s["cid"]}/audit'),
    }
    for k in before:
        assert after[k] == before[k], f'{k} mudou após reinício'
    assert w.call('GET', f'/api/competences/{w.s["cid"]}/integrity')['ok'] is True
    assert w.call('GET', f'/api/presentations/{w.s["pid"]}/integrity')['ok'] is True
    # execuções de coleta vivem só em memória (por desenho): após reiniciar, importar a coleta antiga é 404 explícito
    r = w.api.post('/api/competences/from-run', json={'run_id': w.s['run_c']['run_id'], 'operator': OP})
    assert r.status_code == 404 and error_code(r) == 'metric_run_not_found'


@needs_engine
def test_22b_reexport_after_restart_is_identical(w):
    # jobs/arquivos de exportação são efêmeros (por desenho); a mesma revisão persistida reexporta o MESMO PPTX
    job, f = w.export(w.s['pid'], 'pptx', w.s['final_revision'])
    assert job['sha256'] == w.s['pptx_job']['sha256']
    assert _pptx_inventory(f.content) == w.s['pptx_inv']
    # a chamada de provedor de IA durante a exportação é proibida
    mp = pytest.MonkeyPatch()

    def forbidden(*a, **k):
        raise AssertionError('provedor de interpretação chamado durante a exportação')
    mp.setattr(providers.DeterministicProvider, 'generate', forbidden)
    mp.setattr(providers.ExternalProvider, 'generate', forbidden)
    try:
        job, _ = w.export(w.s['pid'], 'pdf', w.s['final_revision'])
        assert job['status'] == 'done'
    finally:
        mp.undo()


# ---------------------------------------------------------------- 23. segurança do fluxo inteiro

def _unzipped_text(data: bytes) -> str:
    with zipfile.ZipFile(io.BytesIO(data)) as z:
        return ''.join(z.read(n).decode('utf-8', 'ignore') for n in z.namelist()
                       if n.endswith(('.xml', '.rels')))


def test_23_no_secret_or_local_identity_leaks(w):
    secrets = [s for t in ('gamma', 'beta') for s in synthetic_credentials(t)]
    logs = '\n'.join(w.s['logs'])
    assert logs, 'nenhum log capturado (o fluxo deveria registrar eventos)'
    for s in secrets:
        assert s not in logs
    # logs sem textos de interpretação nem valores manuais (F08/F05: só ids, códigos e contagens)
    assert w.s['approved_iv']['text'] not in logs and JUST not in logs
    # banco local: só a referência `env:NOME`, nunca o valor da credencial
    raw = open(w.url.removeprefix('sqlite:///'), 'rb').read()
    assert all(s.encode() not in raw for s in secrets) and b'env:VC_GAMMA' in raw
    # acesso não-loopback e Host forjado são recusados antes de qualquer rota
    assert w.api.get('/api/clients', headers={'Host': 'afrika.example.com'}).status_code == 403
    if 'pptx' not in w.s:  # sem motor de exportação: o restante já foi coberto acima
        return
    user = getpass.getuser()
    pptx_xml = _unzipped_text(w.s['pptx'].content)
    for blob in (pptx_xml, w.s['pdf'].content.decode('latin-1')):
        assert 'VC_GAMMA' not in blob and 'env:' not in blob
        assert all(s not in blob for s in secrets)
        if len(user) >= 3:
            assert user not in blob, 'nome do usuário do sistema operacional no arquivo exportado'
        assert str(w.tmp) not in blob and 'AppData' not in blob and '/home/' not in blob
    # PPTX sem macros, objetos OLE externos nem links remotos
    # Única parte binária: `printerSettings1.bin` herdada byte a byte do template padrão do python-pptx (inerte,
    # igual para qualquer usuário; não é macro nem dado do Afrika) — qualquer outra `.bin` reprova.
    import pptx as pptx_lib
    template = Path(pptx_lib.__file__).parent / 'templates' / 'default.pptx'
    with zipfile.ZipFile(template) as z:
        inherited = {n: z.read(n) for n in z.namelist() if n.endswith('.bin')}
    with zipfile.ZipFile(io.BytesIO(w.s['pptx'].content)) as z:
        names = z.namelist()
        binaries = {n: z.read(n) for n in names if n.endswith('.bin')}
    assert binaries == inherited == {'ppt/printerSettings/printerSettings1.bin':
                                     inherited.get('ppt/printerSettings/printerSettings1.bin')}
    assert not any('vbaProject' in n or 'activeX' in n or 'oleObject' in n
                   or ('embeddings/' in n and not n.endswith('.xlsx')) for n in names)
    assert 'TargetMode="External"' not in pptx_xml
    core = _unzipped_text(w.s['pptx'].content)
    assert 'Steve Canny' not in core and '<dc:creator>Afrika Reporting</dc:creator>' in core
