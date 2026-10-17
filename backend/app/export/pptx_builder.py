"""Construtor do PPTX EDITÁVEL (F07) a partir do layout renderizado + IR aprovada + tokens do tema.

Mapeamento (matriz completa em docs/EXPORT.md):
- texto (títulos, parágrafos, tópicos, rótulos, KPIs, rodapé) -> caixa de texto nativa (fonte, tamanho, cor, entrelinha
  exata, recuos e marcadores do HTML; uma linha => sem quebra automática; várias => quebra na largura do HTML);
- fundos de cartões/faixas/pílulas/filetes/amostras -> formas vetoriais (retângulo, arredondado, elipse ou
  `custGeom`) com preenchimento/alfa/contorno; contornos do tema (caminho SVG) -> `custGeom` com curvas;
- tabelas (comparação, ranking, indicadores) e matriz -> tabela nativa (células multilinha; a matriz emula o espaço
  entre células com bordas na cor do fundo);
- gráficos (barras, linhas, combinado, severidade, rosca) -> gráfico nativo com planilha incorporada (pptx_charts);
- imagem SOMENTE para ativo de marca/tema autorizado (picture), marcada como não editável no manifesto.
Nada é recalculado: textos são os rótulos já formatados da IR; números dos gráficos são os `y` aprovados.
"""
from __future__ import annotations

import io
from dataclasses import dataclass, field
from datetime import UTC, datetime

from lxml import etree
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_CONNECTOR, MSO_SHAPE
from pptx.opc.constants import RELATIONSHIP_TYPE as RT
from pptx.oxml.ns import qn

from app.export import pptx_charts
from app.export import pptx_xml as X
from app.export.fonts import FontBook

GENERATOR = 'afrika-reporting-export/1'
CSS_GENERIC = {'system-ui', '-apple-system', 'BlinkMacSystemFont', 'sans-serif', 'serif', 'monospace', 'cursive',
               'fantasy', 'ui-sans-serif', 'ui-serif', 'ui-monospace', 'ui-rounded'}
# PowerPoint com entrelinha EXATA L põe a 1ª linha de base a 0,75·L do topo da caixa (medido na F07 para Montserrat,
# Poppins e Be Vietnam Pro, 16–107 px, entrelinhas 1,0–1,5 e "normal"); o CSS a põe em meia-entrelinha + ascendente.
PPT_BASELINE = 0.75


@dataclass
class BuildInput:
    view: dict
    layouts: list[dict]
    fontbook: FontBook
    images: dict[str, tuple[bytes, str]] = field(default_factory=dict)  # 'brand:slot'/'image:nome' -> (bytes, tipo)
    rasters: dict[str, bytes] = field(default_factory=dict)  # PNG de reserva para SVG
    meta: dict = field(default_factory=dict)


@dataclass
class Built:
    prs: object
    manifest: list[list[dict]]
    faces: dict[str, set[bool]]
    warnings: list[dict]


def resolve_ir(slide: dict, path: str):
    node = slide
    for part in path.split('.'):
        if not isinstance(node, dict) or part not in node:
            return None
        node = node[part]
    return node


def _ts(view: dict) -> datetime:
    """Data técnica determinística: momento de criação da revisão exportada (UTC, sem fração)."""
    raw = view['revision']['created_at']
    dt = datetime.fromisoformat(raw.replace('Z', '+00:00'))
    if dt.tzinfo is None:
        dt = dt.replace(tzinfo=UTC)
    return dt.astimezone(UTC).replace(microsecond=0, tzinfo=None)


class Builder:
    def __init__(self, inp: BuildInput):
        self.inp, self.view, self.theme = inp, inp.view, inp.view['theme']
        self.faces: dict[str, set[bool]] = {}
        self.warnings: list[dict] = []
        self.manifest: list[list[dict]] = []
        # família CSS sem arquivo empacotado (ex.: pilha do sistema) -> `export_family` declarado pelo tema para o papel
        fonts = self.theme['fonts']
        self.export_family = {fonts[r]['stack'][0]: fonts[r]['export_family'] for r in ('heading', 'body', 'numeric')}
        self.generic_family = fonts['body']['export_family']

    # ------------------------------------------------------------ utilidades
    def face(self, family: str, weight: float) -> tuple[str, bool]:
        book = self.inp.fontbook
        if not book.bundled(family):
            family = self.export_family.get(family) or (self.generic_family if family in CSS_GENERIC else family)
        f = book.face(family, weight)
        self.faces.setdefault(f.typeface, set()).add(f.bold)
        return f.typeface, f.bold

    def _ref(self, competence: str | None) -> dict | None:
        src = self.view['sources']
        if competence and competence == src['primary']['competence']:
            return src['primary']
        return (src.get('history') or {}).get(competence or '')

    def prov_text(self, prov: str | None) -> str | None:
        if not prov:
            return None
        metric, comp, state, origin = (prov.split('|') + ['', '', '', ''])[:4]
        ref = self._ref(comp)
        src = (f'revisão aprovada {ref["revision"]} (aprovação {ref["approval_sequence"]})' if ref
               else 'sem revisão aprovada fixada (lacuna)')
        return (f'Indicador {metric} · competência {comp or "—"} · estado {state or "—"} · origem '
                f'{origin or "—"} · {src}')

    def _record(self, items: list[dict], key: str, kind: str, box: dict, editable: bool = True, **extra) -> None:
        items.append({'key': key, 'kind': kind, 'box': {k: round(v, 2) for k, v in box.items()}, 'editable': editable,
                      **extra})

    # ------------------------------------------------------------ montagem
    def build(self) -> Built:
        prs = Presentation()
        canvas = self.theme['canvas']
        prs.slide_width, prs.slide_height = X.emu(canvas['width']), X.emu(canvas['height'])
        self._theme_part(prs)
        blank = prs.slide_layouts[6]
        slides = self.view['slides']
        if len(slides) != len(self.inp.layouts):
            raise ValueError('layout capturado não corresponde aos slides do documento')
        for vs, lay in zip(slides, self.inp.layouts):
            if lay['slide_id'] != vs['slide_id'] or lay['index'] != vs['index']:
                raise ValueError('ordem dos slides do layout diverge do documento')
            slide = prs.slides.add_slide(blank)
            items: list[dict] = []
            for op in lay['ops']:
                self._op(slide, vs, op, items)
            notes = (vs.get('fields') or {}).get('notes')
            if notes:
                slide.notes_slide.notes_text_frame.text = notes
            self.manifest.append(items)
        self._core(prs)
        return Built(prs, self.manifest, self.faces, self.warnings)

    def _op(self, slide, vs: dict, op: dict, items: list[dict]) -> None:
        kind = op['op']
        shapes = slide.shapes
        if kind == 'background':
            if op.get('fill'):
                fill = slide.background.fill
                fill.solid()
                fill.fore_color.rgb = RGBColor.from_string(op['fill']['rgb'])
            if op.get('image'):
                self.warnings.append({'code': 'background_image_as_picture', 'slide': vs['index'],
                                      'message': 'Imagem de fundo do tema exportada como imagem (ativo do tema)'})
            return
        if kind == 'rect':
            return self._rect(shapes, op, items)
        if kind == 'line':
            return self._line(shapes, op, items)
        if kind == 'arrow':
            return self._arrow(shapes, op, items)
        if kind == 'path':
            return self._path(shapes, op, items)
        if kind == 'text':
            return self._text(shapes, op, items)
        if kind == 'image':
            return self._image(shapes, op, items, vs)
        if kind == 'table':
            return self._table(shapes, op, items)
        if kind == 'chart':
            return self._chart(shapes, op, items, vs)
        raise ValueError(f'operação de layout desconhecida: {kind}')

    # ------------------------------------------------------------ formas
    def _rect(self, shapes, op: dict, items: list[dict]) -> None:
        b = op['box']
        x, y, w, h = b['x'], b['y'], b['w'], b['h']
        r = list(op['radius'])
        line = op.get('border')
        if line:  # contorno do PowerPoint é centrado na borda; a borda CSS fica por dentro da caixa
            i = line['w'] / 2
            x, y, w, h = x + i, y + i, w - 2 * i, h - 2 * i
            r = [max(0.0, v - i) for v in r]
        if w <= 0.1 or h <= 0.1:
            return
        uniform = max(r) - min(r) < 0.05
        if max(r) < 0.05:
            st = MSO_SHAPE.RECTANGLE
        elif uniform and r[0] >= min(w, h) / 2 - 0.05 and abs(w - h) < 0.5:
            st = MSO_SHAPE.OVAL
        elif uniform:
            st = MSO_SHAPE.ROUNDED_RECTANGLE
        else:
            st = MSO_SHAPE.RECTANGLE
        sp = shapes.add_shape(st, X.emu(x), X.emu(y), X.emu(w), X.emu(h))
        if st == MSO_SHAPE.ROUNDED_RECTANGLE:
            sp.adjustments[0] = min(0.5, r[0] / min(w, h))
        if not uniform and max(r) >= 0.05:
            X.cust_geom(sp._element.spPr, X.rounded_rect_segments(w, h, r), w, h)
        X.drop_style(sp._element)
        X.replace_fill(sp._element.spPr, op.get('fill'), op.get('texture', False))
        X.replace_line(sp._element.spPr, line)
        X.set_descr(sp, op['key'], self.prov_text(op.get('prov')))
        self._record(items, op['key'], 'shape', {'x': x, 'y': y, 'w': w, 'h': h})

    def _line(self, shapes, op: dict, items: list[dict]) -> None:
        cx = shapes.add_connector(MSO_CONNECTOR.STRAIGHT, X.emu(op['x1']), X.emu(op['y1']), X.emu(op['x2']),
                                  X.emu(op['y2']))
        X.drop_style(cx._element)
        X.replace_line(cx._element.spPr, {'w': op['w'], 'color': op['color'], 'style': op.get('style', 'solid')})
        X.set_descr(cx, op['key'], None)
        self._record(items, op['key'], 'line', {'x': min(op['x1'], op['x2']), 'y': min(op['y1'], op['y2']),
                                                'w': abs(op['x2'] - op['x1']), 'h': abs(op['y2'] - op['y1'])})

    def _arrow(self, shapes, op: dict, items: list[dict]) -> None:
        b = op['box']
        y = b['y'] + b['h'] / 2
        cx = shapes.add_connector(MSO_CONNECTOR.STRAIGHT, X.emu(b['x']), X.emu(y), X.emu(b['x'] + b['w']), X.emu(y))
        X.drop_style(cx._element)
        X.replace_line(cx._element.spPr, {'w': op['w'], 'color': op['color'], 'style': 'solid'}, tail_arrow=True)
        X.set_descr(cx, op['key'], 'Seta: competência anterior → atual')
        self._record(items, op['key'], 'line', b)

    def _path(self, shapes, op: dict, items: list[dict]) -> None:
        b = op['box']
        sp = shapes.add_shape(MSO_SHAPE.RECTANGLE, X.emu(b['x']), X.emu(b['y']), X.emu(b['w']), X.emu(b['h']))
        X.cust_geom(sp._element.spPr, X.svg_path_segments(op['d']), op['vw'] or b['w'], op['vh'] or b['h'])
        X.drop_style(sp._element)
        X.replace_fill(sp._element.spPr, op['fill'])
        X.replace_line(sp._element.spPr, None)
        X.set_descr(sp, op['key'], 'Forma decorativa do tema (contorno vetorial editável)')
        self._record(items, op['key'], 'shape', b)

    # ------------------------------------------------------------ texto
    def _text(self, shapes, op: dict, items: list[dict]) -> None:
        b = dict(op['box'])
        if op['paras']:  # alinha a 1ª linha de base do PowerPoint à do HTML (ver PPT_BASELINE)
            shift = op.get('base', PPT_BASELINE * op['paras'][0]['lh']) - PPT_BASELINE * op['paras'][0]['lh']
            b['x' if op['vertical'] else 'y'] += shift
        if op['wrap'] and not op['vertical']:
            # folga mínima para diferenças submilimétricas de métrica entre motores (sem mudar as quebras do HTML)
            slack = max(1.0, b['w'] * 0.01)
            align = op['paras'][0]['align'] if op['paras'] else 'l'
            b['x'] -= slack / 2 if align == 'ctr' else slack if align == 'r' else 0
            b['w'] += slack
        tb = shapes.add_textbox(X.emu(b['x']), X.emu(b['y']), X.emu(max(b['w'], 1)), X.emu(max(b['h'], 1)))
        body = tb._element.txBody
        X.body_pr(body, wrap=op['wrap'], vertical=op['vertical'])
        X.write_paragraphs(body, op['paras'], self.face)
        X.set_descr(tb, op['key'], self.prov_text(op.get('prov')))
        self._record(items, op['key'], 'text', op['box'], lines=op['lines'],
                     text=''.join(r['t'] for p in op['paras'] for r in p['runs']))

    # ------------------------------------------------------------ imagens (somente ativos autorizados)
    def _image(self, shapes, op: dict, items: list[dict], vs: dict) -> None:
        key = op['key']
        asset = self.inp.images.get(key)
        if asset is None:
            self.warnings.append({'code': 'image_missing', 'slide': vs['index'], 'key': key,
                                  'message': 'Ativo de imagem indisponível na exportação: encaixe vazio (fallback)'})
            return
        data, media = asset
        png = self.inp.rasters.get(key) if media == 'image/svg+xml' else data
        if png is None:
            self.warnings.append({'code': 'svg_without_fallback', 'slide': vs['index'], 'key': key,
                                  'message': 'Ativo SVG sem imagem de reserva: não exportado'})
            return
        b = op['box']
        nw, nh = op.get('natural') or [0, 0]
        x, y, w, h = b['x'], b['y'], b['w'], b['h']
        crop = None
        if nw and nh and w and h:
            ia, ba = nw / nh, w / h
            if op.get('fit') == 'contain':
                if ia > ba:
                    nh2 = w / ia
                    y, h = y + (h - nh2) / 2, nh2
                else:
                    nw2 = h * ia
                    x, w = x + (w - nw2) / 2, nw2
            elif op.get('fit') == 'cover':
                crop = ((1 - ba / ia) / 2, 0.0) if ia > ba else (0.0, (1 - ia / ba) / 2)
        pic = shapes.add_picture(io.BytesIO(png), X.emu(x), X.emu(y), X.emu(w), X.emu(h))
        if crop:
            pic.crop_left = pic.crop_right = crop[0]
            pic.crop_top = pic.crop_bottom = crop[1]
        if media == 'image/svg+xml':
            self._attach_svg(pic, data)
        X.set_descr(pic, key, f'Ativo visual autorizado ({key.split(":", 1)[-1]}) — imagem (não editável como '
                              'texto/forma; nenhum dado)')
        self._record(items, key, 'picture', {'x': x, 'y': y, 'w': w, 'h': h}, editable=False)

    def _attach_svg(self, pic, svg: bytes) -> None:
        from pptx.opc.package import Part
        from pptx.opc.packuri import PackURI
        slide_part = pic.part
        pkg = slide_part.package
        partname = pkg.next_partname('/ppt/media/image%d.svg')
        part = Part(PackURI(partname), 'image/svg+xml', pkg, svg)
        rid = slide_part.relate_to(part, RT.IMAGE)
        blip = pic._element.find('.//' + qn('a:blip'))
        ext_lst = blip.find(qn('a:extLst'))
        if ext_lst is None:
            ext_lst = X.sub(blip, 'a:extLst')
        ext = X.sub(ext_lst, 'a:ext', uri='{96DAC541-7B7A-43D3-8B79-37D633B846F1}')
        el = etree.SubElement(ext, '{http://schemas.microsoft.com/office/drawing/2016/SVG/main}svgBlip',
                              nsmap={'asvg': 'http://schemas.microsoft.com/office/drawing/2016/SVG/main'})
        el.set(qn('r:embed'), rid)

    # ------------------------------------------------------------ tabelas
    def _table(self, shapes, op: dict, items: list[dict]) -> None:
        rows = [r for r in op['rows'] if r['cells']]
        if not rows:
            return
        ncols = max(len(r['cells']) for r in rows)
        cols = list(op['cols']) + [op['cols'][-1] if op['cols'] else 50.0] * (ncols - len(op['cols']))
        gap = op.get('gap') or 0.0
        box = op['box']
        x0, y0 = box['x'], rows[0]['y']
        widths = [c + gap for c in cols[:ncols]]
        heights = [r['h'] + gap for r in rows]
        if gap:
            x0, y0 = x0 - gap / 2, y0 - gap / 2
        gf = shapes.add_table(len(rows), ncols, X.emu(x0), X.emu(y0), X.emu(sum(widths)), X.emu(sum(heights)))
        tbl = gf.table
        tbl_pr = tbl._tbl.tblPr
        tbl_pr.set('firstRow', '1' if rows[0]['header'] else '0')
        tbl_pr.set('bandRow', '0')
        style = tbl_pr.find(qn('a:tableStyleId'))
        if style is None:
            style = X.sub(tbl_pr, 'a:tableStyleId')
        style.text = '{2D5ABB26-0587-4C30-8999-92F81FD0307C}'  # "Sem estilo, sem grade": formatação só explícita
        for i, w in enumerate(widths):
            tbl.columns[i].width = X.emu(w)
        for j, h in enumerate(heights):
            tbl.rows[j].height = X.emu(h)
        gap_line = {'w': gap, 'color': op.get('gapColor') or {'rgb': 'FFFFFF', 'a': 1}, 'style': 'solid'} if gap else None
        for j, row in enumerate(rows):
            for i in range(ncols):
                cell = tbl.cell(j, i)
                data = row['cells'][i] if i < len(row['cells']) else None
                self._cell(cell, data, gap, gap_line)
        gf.width, gf.height = X.emu(sum(widths)), X.emu(sum(heights))
        X.set_descr(gf, op['key'], 'Tabela nativa editável' + (' (matriz)' if op['kind'] == 'matrix' else ''))
        self._record(items, op['key'], 'table', {'x': x0, 'y': y0, 'w': sum(widths), 'h': sum(heights)},
                     rows=len(rows), cols=ncols,
                     cells=[[''.join(r['t'] for p in (c['paras'] if c else []) for r in p['runs'])
                             for c in row['cells']] for row in rows])

    def _cell(self, cell, data: dict | None, gap: float, gap_line: dict | None) -> None:
        tc = cell._tc
        body = tc.find(qn('a:txBody'))
        bp = body.find(qn('a:bodyPr'))
        for k in list(bp.attrib):
            del bp.attrib[k]
        X.write_paragraphs(body, (data or {}).get('paras') or [], self.face)
        old = tc.find(qn('a:tcPr'))
        if old is not None:
            tc.remove(old)
        pad = (data or {}).get('pad') or [0, 0, 0, 0]
        extra = gap / 2
        valign = (data or {}).get('valign', 't')
        shift = 0.0
        if data and data.get('paras'):  # mesma correção de linha de base das caixas de texto, via margens
            lh = data['paras'][0]['lh']
            shift = data.get('base', PPT_BASELINE * lh) - PPT_BASELINE * lh
        top = pad[0] + extra + (shift if valign in ('t', 'ctr') else 0)
        bottom = pad[2] + extra - (shift if valign in ('ctr', 'b') else 0)
        # O PowerPoint AUMENTA a linha quando margem sup. + altura natural do texto (métrica da fonte, maior que a
        # entrelinha exata) + margem inf. passa da altura — e empurra as linhas seguintes (cabeçalho da matriz: +5 px).
        # A margem do lado oposto à âncora não move o texto (no centro, só a diferença entre elas move): reduzi-la
        # mantém a posição e libera a altura.
        if valign == 't':
            bottom = 0.0
        elif valign == 'b':
            top = 0.0
        else:
            cut = max(0.0, min(top, bottom))
            top, bottom = top - cut, bottom - cut
        tcpr = X.sub(tc, 'a:tcPr', marL=X.emu(pad[3] + extra), marR=X.emu(pad[1] + extra),
                     marT=X.emu(max(0.0, top)), marB=X.emu(max(0.0, bottom)), anchor=valign)
        borders = (data or {}).get('borders') or [None, None, None, None]
        # ordem do esquema: lnL, lnR, lnT, lnB, …, preenchimento
        for tag, b in (('a:lnL', borders[3]), ('a:lnR', borders[1]), ('a:lnT', borders[0]), ('a:lnB', borders[2])):
            line = gap_line or b
            ln = X.sub(tcpr, tag)
            if line:
                ln.set('w', str(X.emu(line['w'])))
                ln.set('cap', 'flat')
                ln.set('cmpd', 'sng')
                ln.set('algn', 'ctr')
                X.solid_fill(ln, line['color'])
                X.sub(ln, 'a:prstDash', val=X.DASH.get(line.get('style', 'solid'), 'solid'))
            else:
                ln.set('w', '0')
                X.sub(ln, 'a:noFill')
        fill = (data or {}).get('fill')
        if fill:
            X.solid_fill(tcpr, fill)
        else:
            X.sub(tcpr, 'a:noFill')

    # ------------------------------------------------------------ gráficos nativos
    def _chart(self, shapes, op: dict, items: list[dict], vs: dict) -> None:
        ir = resolve_ir(vs, op['ir'])
        if not isinstance(ir, dict) or ir.get('available') is False:
            raise ValueError(f'gráfico sem IR disponível em {op["ir"]} (slide {vs["index"]})')
        if op['variant'] == 'donut':
            gf = pptx_charts.add_donut(shapes, op, ir, self.theme, self.face)
            kind = 'doughnut'
        elif op['variant'] == 'severity':
            gf = pptx_charts.add_severity_chart(shapes, op, ir, self.theme, self.face)
            kind = 'bar_horizontal'
        else:
            gf = pptx_charts.add_plot_chart(shapes, op, ir, self.theme, self.face)
            kind = {'line': 'line', 'combo': 'combo'}.get(ir['kind'], 'bar')
        refs = sorted({f'{r["metric_id"]}@{r["competence"]} (r{r["revision"]}/a{r["approval_sequence"]})'
                       for r in ir.get('data_refs', []) if r.get('revision')})
        gaps = ir.get('gaps') or []
        descr = (f'Gráfico nativo editável ({kind}) com dados incorporados. Séries: '
                 f'{", ".join(s["label"] for s in ir.get("series", []))}. Lacunas: {", ".join(gaps) or "nenhuma"} '
                 f'(células vazias, sem interpolação). Fontes: {"; ".join(refs[:24])}')
        X.set_descr(gf, op['key'], descr)
        self._record(items, op['key'], 'chart', op['box'], chart=kind, ir=op['ir'],
                     values=pptx_charts.chart_values(gf))
        for ov in op.get('overlays') or []:  # marcadores de lacuna, rótulo de unidade, total da rosca (vetoriais)
            self._overlay(shapes, ov, items)

    def _overlay(self, shapes, op: dict, items: list[dict]) -> None:
        {'rect': self._rect, 'line': self._line, 'text': self._text}[op['op']](shapes, op, items)

    # ------------------------------------------------------------ tema e propriedades
    def _theme_part(self, prs) -> None:
        """Paleta e fontes do tema do PowerPoint = tokens do tema (comodidade de edição; formas usam cor explícita)."""
        part = prs.slide_master.part.part_related_by(RT.THEME)
        root = etree.fromstring(part.blob)
        a = 'http://schemas.openxmlformats.org/drawingml/2006/main'
        c = self.theme['colors']

        def pick(*tokens: str) -> str:
            for t in tokens:
                if c.get(t):
                    return c[t].lstrip('#')[:6].upper()
            return '000000'

        scheme = {'dk1': pick('text.primary'), 'lt1': pick('surface'), 'dk2': pick('title', 'brand.deep', 'text.primary'),
                  'lt2': pick('band', 'surface'), 'accent1': pick('series.1'), 'accent2': pick('series.2'),
                  'accent3': pick('series.3'), 'accent4': pick('series.4'),
                  'accent5': pick('severity.VERY_HIGH', 'accent'), 'accent6': pick('severity.LOW', 'accent'),
                  'hlink': pick('accent'), 'folHlink': pick('series.2')}
        clr = root.find(f'.//{{{a}}}clrScheme')
        clr.set('name', f'{self.theme["id"]}@{self.theme["version"]}')
        for name, val in scheme.items():
            el = clr.find(f'{{{a}}}{name}')
            for child in list(el):
                el.remove(child)
            etree.SubElement(el, f'{{{a}}}srgbClr', val=val)
        fonts = self.theme['fonts']
        heading = self.face(fonts['heading']['stack'][0], self.theme['weights'].get('title', 700))[0]
        body = self.face(fonts['body']['stack'][0], 400)[0]
        fs = root.find(f'.//{{{a}}}fontScheme')
        fs.set('name', f'{self.theme["id"]}@{self.theme["version"]}')
        fs.find(f'{{{a}}}majorFont/{{{a}}}latin').set('typeface', heading)
        fs.find(f'{{{a}}}minorFont/{{{a}}}latin').set('typeface', body)
        part._blob = etree.tostring(root, xml_declaration=True, encoding='UTF-8', standalone=True)

    def _core(self, prs) -> None:
        v = self.view
        cp = prs.core_properties
        ts = _ts(v)
        tmpl = v['template']
        cp.title = v['presentation']['title'][:250]
        cp.subject = f'{v["owner"]["name"]} — competência {v["competence"]["competence"]}'[:250]
        cp.author = 'Afrika Reporting'
        cp.last_modified_by = 'Afrika Reporting — exportação F07'
        cp.keywords = (f'template {tmpl["key"]}@{tmpl["version"]}; tema {self.theme["id"]}@{self.theme["version"]}; '
                       f'revisão {v["revision"]["number"]}')
        cp.comments = ('Exportado da revisão PERSISTIDA da apresentação; valores somente das revisões aprovadas '
                       'fixadas (sem recálculo); lacunas permanecem vazias. Proveniência completa nas propriedades '
                       'personalizadas e no texto alternativo dos objetos.')
        cp.category = 'Relatório AppSec mensal'
        cp.revision = int(v['revision']['number'])
        cp.created = ts
        cp.modified = ts
        cp.language = 'pt-BR'
        cp.identifier = f'{v["presentation"]["id"]}#r{v["revision"]["number"]}'
        cp.version = GENERATOR
        cp.content_status = 'Exportado'


def build(inp: BuildInput) -> Built:
    return Builder(inp).build()
