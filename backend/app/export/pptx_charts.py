"""Gráficos NATIVOS do PowerPoint (F07) — planilha incorporada com os valores APROVADOS da IR, nunca imagem.

- Valores: `y` exato da IR para pontos `VALUE`; lacuna/ausente/incompatível = célula VAZIA (`None`) e
  `dispBlanksAs="gap"` (linha quebrada, sem barra) — nunca 0, nunca interpolação.
- Rótulos de dados: texto pt-BR já formatado pela IR (milhar "." e decimal ","), independente do idioma do PowerPoint
  de quem abre (rótulo editável; ao editar dados, atualizar ou remover o rótulo personalizado).
- Geometria: a área interna de plotagem é posicionada na MESMA caixa do HTML (`data-plot-area` medido no DOM); eixo Y
  de 0 ao mesmo máximo "redondo" do HTML; largura das barras pela mesma razão do tema; cores = tokens do tema.
- Tipos: barras (colunas), linhas com marcadores, combinado (1ª série em barras + demais em linhas no MESMO eixo),
  severidade (barras horizontais, cor por ponto) e rosca (cor por severidade, filete na cor da superfície).
"""
from __future__ import annotations

import math

from lxml import etree
from pptx.chart.data import CategoryChartData
from pptx.dml.color import RGBColor
from pptx.enum.dml import MSO_PATTERN
from pptx.enum.chart import XL_CHART_TYPE, XL_LABEL_POSITION, XL_MARKER_STYLE, XL_TICK_LABEL_POSITION, XL_TICK_MARK
from pptx.oxml.ns import qn
from pptx.util import Emu, Pt

from app.export.pptx_xml import emu, hundredths_pt, sub

NUMBER_FORMATS = {'count': '#,##0', 'by_severity': '#,##0', 'breakdown': '#,##0', 'ranking': '#,##0',
                  'decimal': '#,##0.00', 'percentage': '#,##0.0"%"'}


def nice_max(values: list[float]) -> float:
    """Mesmo máximo de eixo do HTML (`niceMax` em SlideView.tsx): 1/2/5/10 × 10^k acima do maior valor."""
    top = max([0.0, *values])
    if top == 0:
        return 1.0
    exp = 10 ** math.floor(math.log10(top))
    f = top / exp
    return (1 if f <= 1 else 2 if f <= 2 else 5 if f <= 5 else 10) * exp


def axis_units(top: float, ticks: int, integer: bool) -> tuple[float, float | None]:
    """(intervalo principal com rótulo, intervalo secundário só de grade) iguais ao HTML: grade a cada máximo/divisões;
    em eixo inteiro, o HTML só rotula as divisões inteiras => rótulo no menor múltiplo inteiro da divisão e grade
    secundária na divisão original (mesmas linhas e mesmos rótulos)."""
    unit = top / max(1, ticks)
    if not integer or float(unit).is_integer():
        return unit, None
    for m in range(2, ticks + 1):
        if abs(unit * m - round(unit * m)) < 1e-9 and round(top / (unit * m), 9).is_integer():
            return unit * m, unit
    return top, unit


def values_of(series: dict) -> list[float | None]:
    return [p['y'] if p.get('state') == 'VALUE' and p.get('y') is not None else None for p in series['points']]


def _rgb(theme: dict, token: str | None, fallback: str = 'series.1') -> RGBColor:
    hexv = theme['colors'].get(token or fallback) or theme['colors'].get(fallback) or '#000000'
    return RGBColor.from_string(hexv.lstrip('#')[:6].upper())


def _hex(theme: dict, token: str) -> str:
    return (theme['colors'].get(token) or '#000000').lstrip('#')[:6].upper()


def _pt(px: float) -> Pt:
    return Pt(px * 0.75)


def _font(run_font, theme: dict, family: str, size_px: float, color_token: str, weight: int, face_of) -> None:
    typeface, bold = face_of(family, weight)
    run_font.name = typeface
    run_font.size = _pt(size_px)
    run_font.bold = bold
    run_font.color.rgb = RGBColor.from_string(_hex(theme, color_token))


def _family(theme: dict, role: str) -> str:
    return theme['fonts'][role]['stack'][0]


def _no_fill_space(chart) -> None:
    """Fundo/contorno transparentes na área do gráfico e na de plotagem (o slide aparece por baixo, como no HTML)."""
    cs = chart._chartSpace
    if cs.find(qn('c:roundedCorners')) is None:
        rc = etree.Element(qn('c:roundedCorners'), val='0')
        anchor = cs.find(qn('c:date1904'))
        (anchor.addnext(rc) if anchor is not None else cs.insert(0, rc))
    sp = cs.find(qn('c:spPr'))
    if sp is None:
        sp = etree.Element(qn('c:spPr'))
        cs.find(qn('c:chart')).addnext(sp)
    for child in list(sp):
        sp.remove(child)
    sub(sp, 'a:noFill')
    sub(sub(sp, 'a:ln'), 'a:noFill')
    plot_area = cs.chart.plotArea
    psp = plot_area.find(qn('c:spPr'))
    if psp is None:
        psp = etree.SubElement(plot_area, qn('c:spPr'))
    for child in list(psp):
        psp.remove(child)
    sub(psp, 'a:noFill')
    sub(sub(psp, 'a:ln'), 'a:noFill')
    chart_el = cs.chart
    blanks = chart_el.find(qn('c:dispBlanksAs'))
    if blanks is None:
        blanks = etree.Element(qn('c:dispBlanksAs'))
        pvo = chart_el.find(qn('c:plotVisOnly'))
        (pvo.addnext(blanks) if pvo is not None else chart_el.append(blanks))
    blanks.set('val', 'gap')  # lacuna = ausência (linha quebrada, sem barra) — nunca zero


def _manual_layout(chart, frame: dict, plot: dict | None) -> None:
    if not plot or frame['w'] <= 0 or frame['h'] <= 0:
        return
    pa = chart._chartSpace.chart.plotArea
    old = pa.find(qn('c:layout'))
    lay = etree.Element(qn('c:layout'))
    ml = sub(lay, 'c:manualLayout')
    sub(ml, 'c:layoutTarget', val='inner')
    sub(ml, 'c:xMode', val='edge')
    sub(ml, 'c:yMode', val='edge')
    clamp = lambda v: max(0.0, min(1.0, v))  # noqa: E731
    sub(ml, 'c:x', val=f'{clamp((plot["x"] - frame["x"]) / frame["w"]):.5f}')
    sub(ml, 'c:y', val=f'{clamp((plot["y"] - frame["y"]) / frame["h"]):.5f}')
    sub(ml, 'c:w', val=f'{clamp(plot["w"] / frame["w"]):.5f}')
    sub(ml, 'c:h', val=f'{clamp(plot["h"] / frame["h"]):.5f}')
    if old is not None:
        pa.replace(old, lay)
    else:
        pa.insert(0, lay)


def _tick_skip(category_axis, n: int) -> None:
    if n <= 1:
        return
    cat = category_axis._element
    skip = cat.find(qn('c:tickLblSkip'))
    if skip is None:
        skip = etree.Element(qn('c:tickLblSkip'))
        anchor = None
        for tag in ('c:lblOffset', 'c:lblAlgn', 'c:auto', 'c:crosses', 'c:crossesAt', 'c:crossAx'):
            anchor = cat.find(qn(tag))
            if anchor is not None:
                break
        (anchor.addnext(skip) if anchor is not None else cat.append(skip))
    skip.set('val', str(n))


def _label_every(drawn: list[str]) -> int:
    """Rarefação de rótulos do eixo de meses como DESENHADA no HTML (índices com rótulo)."""
    shown = [i for i, t in enumerate(drawn) if t]
    if len(shown) < 2:
        return 1
    return max(1, shown[1] - shown[0])


def _axes(chart, theme: dict, ir: dict, top: float, face_of, *, mid_cat: bool) -> None:
    c = theme['charts']
    va, ca = chart.value_axis, chart.category_axis
    va.minimum_scale = 0
    va.maximum_scale = top
    major, minor = axis_units(top, int(c.get('y_ticks', 4)), bool(ir.get('y_axis', {}).get('integer', True)))
    va.major_unit = major
    va.has_major_gridlines = True
    gl = va.major_gridlines.format.line
    gl.color.rgb = RGBColor.from_string(_hex(theme, 'grid'))
    gl.width = Emu(emu(c.get('grid_width', 1)))
    if minor:  # python-pptx não formata a grade secundária: XML direto (mesma cor/espessura da principal)
        va.minor_unit = minor
        va.has_minor_gridlines = True
        mg = va._element.find(qn('c:minorGridlines'))
        sp = sub(mg, 'c:spPr')
        ln = sub(sp, 'a:ln', w=emu(c.get('grid_width', 1)))
        sub(sub(ln, 'a:solidFill'), 'a:srgbClr', val=_hex(theme, 'grid'))
    va.format.line.fill.background()
    va.major_tick_mark = XL_TICK_MARK.NONE
    va.minor_tick_mark = XL_TICK_MARK.NONE
    va.tick_label_position = XL_TICK_LABEL_POSITION.LOW
    va.tick_labels.number_format = '#,##0' if ir.get('y_axis', {}).get('integer', True) else '#,##0.##'
    va.tick_labels.number_format_is_linked = False
    _font(va.tick_labels.font, theme, _family(theme, 'body'), theme['sizes']['axis'], 'text.muted', 400, face_of)
    ca.format.line.color.rgb = RGBColor.from_string(_hex(theme, 'axis'))
    ca.format.line.width = Emu(emu(c.get('grid_width', 1)))
    ca.major_tick_mark = XL_TICK_MARK.NONE
    ca.minor_tick_mark = XL_TICK_MARK.NONE
    ca.tick_label_position = XL_TICK_LABEL_POSITION.LOW
    _font(ca.tick_labels.font, theme, _family(theme, 'body'), theme['sizes']['axis'], 'text.muted', 400, face_of)
    va_el = va._element
    cb = va_el.find(qn('c:crossBetween'))
    if cb is None:
        cb = etree.Element(qn('c:crossBetween'))
        anchor = va_el.find(qn('c:crosses'))
        if anchor is None:
            anchor = va_el.find(qn('c:crossAx'))
        anchor.addnext(cb)
    # linhas: pontos nas bordas da área (como o HTML: 1º mês na borda esquerda); barras/combinado: centro da faixa
    cb.set('val', 'midCat' if mid_cat else 'between')


def _labels(series, points: list[dict], theme: dict, face_of, position, *, weight_token: str = 'value_label') -> None:
    """Rótulo de dados personalizado por ponto COM valor (texto pt-BR da IR); lacunas sem rótulo (marcador à parte)."""
    weight = int(theme['weights'].get(weight_token, 600))
    for j, p in enumerate(points):
        if p.get('state') != 'VALUE' or p.get('y') is None:
            continue
        dl = series.points[j].data_label
        tf = dl.text_frame
        tf.text = p['label']
        _font(tf.paragraphs[0].runs[0].font, theme, _family(theme, 'numeric'), theme['sizes']['value_label'],
              'text.primary', weight, face_of)
        dl.position = position


def _bar_gap(theme: dict, op: dict, n: int) -> int:
    c = theme['charts']
    pad = c['plot_padding']
    vw = op.get('vb', [0, 0])[0] or c['plot_width']
    iw = max(1.0, vw - pad['left'] - pad['right'])
    step = iw / max(1, n)
    bw = min(c['bar_max_width'], step * c['bar_width_ratio'])
    return max(0, min(500, int(round((step / max(bw, 0.1) - 1) * 100))))


def _data(ir: dict, series: list[dict], value_type: str) -> CategoryChartData:
    fmt = NUMBER_FORMATS.get(value_type, '#,##0')
    cd = CategoryChartData(number_format=fmt)
    cd.categories = list(ir['categories'])
    for s in series:
        cd.add_series(s['label'], values_of(s), number_format=fmt)
    return cd


def _value_type(ir: dict) -> str:
    nf = ir.get('number_format') or {}
    if nf.get('suffix') == '%':
        return 'percentage'
    return 'decimal' if nf.get('decimals') == 2 else 'count'


def add_plot_chart(shapes, op: dict, ir: dict, theme: dict, face_of):
    """Barras / linhas / combinado (SVG do HTML)."""
    kind = ir['kind']
    box = op['box']
    series = ir['series']
    vt = _value_type(ir)
    if kind == 'line':
        ctype = XL_CHART_TYPE.LINE_MARKERS
    else:
        ctype = XL_CHART_TYPE.COLUMN_CLUSTERED
    data = _data(ir, series, vt)
    gf = shapes.add_chart(ctype, emu(box['x']), emu(box['y']), emu(box['w']), emu(box['h']), data)
    chart = gf.chart
    chart.has_legend = False
    chart.has_title = False
    chart.font.size = _pt(theme['sizes']['axis'])
    chart.font.name = face_of(_family(theme, 'body'), 400)[0]
    allv = [v for s in series for v in values_of(s) if v is not None] if kind in ('line', 'combo') \
        else [v for v in values_of(series[0]) if v is not None]
    top = nice_max(allv)
    _axes(chart, theme, ir, top, face_of, mid_cat=kind == 'line')
    _tick_skip(chart.category_axis, _label_every(op.get('categories') or []))
    c = theme['charts']
    plot = chart.plots[0]
    if kind in ('bar', 'combo'):
        plot.gap_width = _bar_gap(theme, op, len(ir['categories']))
        plot.overlap = 0
        bar = plot.series[0]
        bar.invert_if_negative = False
        bar.format.fill.solid()
        bar.format.fill.fore_color.rgb = _rgb(theme, series[0].get('color'))
        bar.format.line.fill.background()
        for j, p in enumerate(series[0]['points']):  # cor por ponto quando a IR define (ex.: severidade em barras)
            if p.get('color') and p.get('state') == 'VALUE':
                pt = bar.points[j]
                pt.format.fill.solid()
                pt.format.fill.fore_color.rgb = _rgb(theme, p['color'])
        if ir.get('show_values', True) is not False:
            _labels(bar, series[0]['points'], theme, face_of, XL_LABEL_POSITION.OUTSIDE_END)
    if kind == 'combo' and len(series) > 1:
        _to_line_chart(chart, theme, c)
        plot_line = chart.plots[1]
        for s, ser in zip(series[1:], plot_line.series):
            _style_line(ser, theme, c, s)
            if ir.get('show_values', True) is not False:
                _labels(ser, s['points'], theme, face_of, XL_LABEL_POSITION.ABOVE)
    if kind == 'line':
        for s, ser in zip(series, plot.series):
            _style_line(ser, theme, c, s)
            if ir.get('show_values', True) is not False:
                _labels(ser, s['points'], theme, face_of, XL_LABEL_POSITION.ABOVE)
    _no_fill_space(chart)
    _manual_layout(chart, box, op.get('plot'))
    return gf


def _style_line(ser, theme: dict, c: dict, s: dict) -> None:
    color = _rgb(theme, s.get('color'))
    ser.smooth = False
    ser.format.line.color.rgb = color
    ser.format.line.width = Emu(emu(c['line_width']))
    ser.marker.style = XL_MARKER_STYLE.CIRCLE
    ser.marker.size = max(2, min(72, int(round(c['point_radius'] * 2 * 0.75))))
    ser.marker.format.fill.solid()
    ser.marker.format.fill.fore_color.rgb = color
    ser.marker.format.line.color.rgb = RGBColor.from_string(_hex(theme, 'surface'))
    ser.marker.format.line.width = Emu(emu(c.get('point_stroke', 2)))


def _to_line_chart(chart, theme: dict, c: dict) -> None:
    """Combinado: move as séries 2..n do barChart para um lineChart com os MESMOS eixos (sem eixo secundário)."""
    pa = chart._chartSpace.chart.plotArea
    bar = pa.find(qn('c:barChart'))
    sers = bar.findall(qn('c:ser'))[1:]
    line = etree.Element(qn('c:lineChart'))
    sub(line, 'c:grouping', val='standard')
    sub(line, 'c:varyColors', val='0')
    for ser in sers:
        bar.remove(ser)
        inv = ser.find(qn('c:invertIfNegative'))
        if inv is not None:
            ser.remove(inv)
        line.append(ser)
        sub(ser, 'c:smooth', val='0')
    sub(line, 'c:marker', val='1')
    for ax in bar.findall(qn('c:axId')):
        sub(line, 'c:axId', val=ax.get('val'))
    bar.addnext(line)
    # ordem do esquema dentro de c:ser (linha): idx, order, tx, spPr, marker, dPt, dLbls, …, cat, val, smooth
    for ser in line.findall(qn('c:ser')):
        smooth = ser.find(qn('c:smooth'))
        ser.remove(smooth)
        ser.append(smooth)


def add_severity_chart(shapes, op: dict, ir: dict, theme: dict, face_of):
    """Distribuição por severidade (linhas do HTML) -> barras HORIZONTAIS nativas, cor por severidade."""
    box = op['box']
    s = ir['series'][0]
    data = _data(ir, [s], 'count')
    gf = shapes.add_chart(XL_CHART_TYPE.BAR_CLUSTERED, emu(box['x']), emu(box['y']), emu(box['w']), emu(box['h']), data)
    chart = gf.chart
    chart.has_legend = False
    chart.has_title = False
    c = theme['charts']
    plot = chart.plots[0]
    rows = max(1, int(op.get('rows') or len(ir['categories'])))
    pitch = (op['plot']['h'] / rows) if op.get('plot') else c['severity_bar_height'] * 1.5
    plot.gap_width = max(0, min(500, int(round((pitch / c['severity_bar_height'] - 1) * 100))))
    plot.overlap = 0
    ser = plot.series[0]
    ser.invert_if_negative = False
    ser.format.line.fill.background()
    for j, p in enumerate(s['points']):
        pt = ser.points[j]
        token = p.get('color') or 'series.1'
        if token == 'severity.UNKNOWN':  # textura da "Desconhecida" (identidade nunca só por cor)
            pt.format.fill.patterned()
            pt.format.fill.pattern = MSO_PATTERN.DARK_UPWARD_DIAGONAL
            pt.format.fill.fore_color.rgb = RGBColor.from_string(_hex(theme, 'surface'))
            pt.format.fill.back_color.rgb = _rgb(theme, token)
        else:
            pt.format.fill.solid()
            pt.format.fill.fore_color.rgb = _rgb(theme, token)
    va, ca = chart.value_axis, chart.category_axis
    va.minimum_scale = 0
    va.maximum_scale = max([1.0, *[v for v in values_of(s) if v is not None]])
    va.has_major_gridlines = False
    va.visible = False
    ca.reverse_order = True  # 1ª severidade (Muito alta/Crítica) no topo, como no HTML
    ca.format.line.fill.background()
    ca.major_tick_mark = XL_TICK_MARK.NONE
    _font(ca.tick_labels.font, theme, _family(theme, 'body'), theme['sizes']['severity'], 'text.primary', 400, face_of)
    if ir.get('show_values', True) is not False:
        _labels(ser, s['points'], theme, face_of, XL_LABEL_POSITION.OUTSIDE_END, weight_token='emphasis')
    _no_fill_space(chart)
    _manual_layout(chart, box, op.get('plot'))
    return gf


def add_donut(shapes, op: dict, ir: dict, theme: dict, face_of):
    """Rosca por severidade: ângulos a partir dos números aprovados; sem percentuais calculados nem rótulos."""
    box = op['box']
    s = ir['series'][0]
    data = _data(ir, [s], 'count')
    gf = shapes.add_chart(XL_CHART_TYPE.DOUGHNUT, emu(box['x']), emu(box['y']), emu(box['w']), emu(box['h']), data)
    chart = gf.chart
    chart.has_legend = False
    chart.has_title = False
    comp = theme['components']
    dn = chart._chartSpace.chart.plotArea.find(qn('c:doughnutChart'))
    vary = dn.find(qn('c:varyColors'))
    if vary is not None:
        vary.set('val', '1')
    first = dn.find(qn('c:firstSliceAng'))
    if first is None:
        first = sub(dn, 'c:firstSliceAng')
    first.set('val', '0')
    hole = dn.find(qn('c:holeSize'))
    if hole is None:
        hole = sub(dn, 'c:holeSize')
    hole.set('val', str(max(10, min(90, int(round(comp.get('donut_hole', 0.6) * 100))))))
    # c:firstSliceAng precede c:holeSize no esquema
    dn.remove(first)
    dn.remove(hole)
    dn.append(first)
    dn.append(hole)
    ser = chart.plots[0].series[0]
    sep = comp.get('donut_separator') or 0
    if sep:
        ser.format.line.color.rgb = RGBColor.from_string(_hex(theme, 'surface'))
        ser.format.line.width = Emu(emu(sep))
    else:
        ser.format.line.fill.background()
    for j, p in enumerate(s['points']):
        pt = ser.points[j]
        pt.format.fill.solid()
        pt.format.fill.fore_color.rgb = _rgb(theme, p.get('color') or 'series.1')
    _no_fill_space(chart)
    # o HTML desenha a rosca com raio 100 num viewBox de 220 (margem de 10 em volta)
    m = box['w'] * 10 / 220
    _manual_layout(chart, box, {'x': box['x'] + m, 'y': box['y'] + m, 'w': box['w'] - 2 * m, 'h': box['h'] - 2 * m})
    return gf


def chart_values(gf) -> list[list[float | None]]:
    """Valores gravados no XML do gráfico (cache), por série, com None nas lacunas — para validação."""
    out = []
    for ser in gf.chart._chartSpace.findall('.//' + qn('c:ser')):
        cache = ser.find('.//' + qn('c:val')).find('.//' + qn('c:numCache'))
        count = int(cache.find(qn('c:ptCount')).get('val'))
        vals: list[float | None] = [None] * count
        for pt in cache.findall(qn('c:pt')):
            vals[int(pt.get('idx'))] = float(pt.find(qn('c:v')).text)
        out.append(vals)
    return out


__all__ = ['add_plot_chart', 'add_severity_chart', 'add_donut', 'chart_values', 'nice_max', 'axis_units',
           'values_of', 'hundredths_pt']
